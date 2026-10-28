import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Color Palette (Modern Dark / Cyber Indigo Theme) ──────────────────────────
BG_COLOR         = RGBColor(10, 13, 22)      # #0A0D16 - Deep midnight obsidian
CARD_BG          = RGBColor(18, 24, 40)      # #121828 - Elevated dark card
CARD_BG_HOVER    = RGBColor(24, 32, 54)      # #182036 - Lighter surface
CARD_BORDER      = RGBColor(38, 52, 84)      # #263454 - Subtle card border
ACCENT_PRIMARY   = RGBColor(99, 102, 241)    # #6366F1 - Electric Indigo
ACCENT_CYAN      = RGBColor(14, 165, 233)    # #0EA5E9 - Cyber Blue
ACCENT_EMERALD   = RGBColor(16, 185, 129)    # #10B981 - Success Green
ACCENT_ROSE      = RGBColor(244, 63, 94)     # #F43F5E - Error / Alert Rose
ACCENT_AMBER     = RGBColor(245, 158, 11)    # #F59E0B - Warning Amber
ACCENT_PURPLE    = RGBColor(168, 85, 247)    # #A855F7 - Modern Violet

TEXT_WHITE       = RGBColor(248, 250, 252)   # #F8FAFC - Pure bright text
TEXT_MUTED       = RGBColor(148, 163, 184)   # #94A3B8 - Slate body text
TEXT_DIM         = RGBColor(100, 116, 139)   # #64748B - Subtitle / footnote text
BADGE_BG         = RGBColor(30, 41, 69)      # #1E2945 - Pill badge fill

FONT_TITLE = "Segoe UI"
FONT_BODY  = "Segoe UI"
FONT_MONO  = "Consolas"

def create_blank_slide(prs):
    """Creates a blank slide with the custom deep midnight background."""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()
    return slide

def add_header(slide, title_text, category_badge="DEVASSIST AI PLATFORM", slide_number=None):
    """Standardized modern header with glowing category badge, title, and bottom divider."""
    # Category badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(2.8), Inches(0.32))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BADGE_BG
    badge.line.color.rgb = ACCENT_PRIMARY
    badge.line.width = Pt(1)
    tf = badge.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.1)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.04)
    p = tf.paragraphs[0]
    p.text = category_badge
    p.font.name = FONT_BODY
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.alignment = PP_ALIGN.CENTER

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.0), Inches(0.6))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = 0
    tf_t.margin_top = 0
    tf_t.margin_right = 0
    tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = FONT_TITLE
    p_t.font.size = Pt(23)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_WHITE

    # Slide Number pill (top right)
    if slide_number:
        num_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.933), Inches(0.35), Inches(0.6), Inches(0.32))
        num_box.fill.solid()
        num_box.fill.fore_color.rgb = BADGE_BG
        num_box.line.color.rgb = CARD_BORDER
        tf_n = num_box.text_frame
        tf_n.margin_top = Inches(0.04)
        pn = tf_n.paragraphs[0]
        pn.text = str(slide_number).zfill(2)
        pn.font.name = FONT_MONO
        pn.font.size = Pt(10)
        pn.font.bold = True
        pn.font.color.rgb = TEXT_MUTED
        pn.alignment = PP_ALIGN.CENTER

    # Thin glowing accent divider line
    div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.38), Inches(11.733), Inches(0.02))
    div.fill.solid()
    div.fill.fore_color.rgb = CARD_BORDER
    div.line.fill.background()

def create_card_with_text(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, accent_bar=None, pad_x=0.22, pad_y=0.18):
    """
    Creates an empty card container and returns a clean, single text frame on top.
    By putting all text in this SINGLE text frame, text elements can NEVER overlap or collide.
    """
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)

    top_offset = pad_y
    if accent_bar:
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent_bar
        bar.line.fill.background()
        top_offset += 0.06

    tb = slide.shapes.add_textbox(
        left + Inches(pad_x),
        top + Inches(top_offset),
        width - Inches(pad_x * 2),
        height - Inches(top_offset + pad_y)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_top = 0
    tf.margin_right = 0
    tf.margin_bottom = 0
    return tf

# ── Main Presentation Builder ─────────────────────────────────────────────────

def build_presentation(out_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # =========================================================================
    # SLIDE 1: Title Slide (Hero)
    # =========================================================================
    s1 = create_blank_slide(prs)

    # Hero Background Box
    hero = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(0.8), Inches(10.933), Inches(4.2))
    hero.fill.solid()
    hero.fill.fore_color.rgb = CARD_BG
    hero.line.color.rgb = ACCENT_PRIMARY
    hero.line.width = Pt(1.5)

    # Top Cyan Strip
    strip = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(0.8), Inches(10.933), Inches(0.08))
    strip.fill.solid()
    strip.fill.fore_color.rgb = ACCENT_CYAN
    strip.line.fill.background()

    # Single text frame for hero content to eliminate any vertical collisions
    h_tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.2), Inches(10.333), Inches(3.6))
    h_tf = h_tb.text_frame
    h_tf.word_wrap = True
    h_tf.margin_left = 0
    h_tf.margin_top = 0
    h_tf.margin_right = 0
    h_tf.margin_bottom = 0

    p_pill = h_tf.paragraphs[0]
    p_pill.text = "[ AUTONOMOUS CODE REVIEW & ARCHITECTURAL AGENT ]"
    p_pill.font.name = FONT_BODY
    p_pill.font.size = Pt(12)
    p_pill.font.bold = True
    p_pill.font.color.rgb = ACCENT_CYAN
    p_pill.alignment = PP_ALIGN.CENTER
    p_pill.space_after = Pt(12)

    p_title = h_tf.add_paragraph()
    p_title.text = "DevAssist AI"
    p_title.font.name = FONT_TITLE
    p_title.font.size = Pt(46)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE
    p_title.alignment = PP_ALIGN.CENTER
    p_title.space_after = Pt(10)

    p_sub = h_tf.add_paragraph()
    p_sub.text = "Next-Generation Multi-LLM Review Routing, AST-Aware CodeGraph & Hybrid RAG"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(16)
    p_sub.font.bold = False
    p_sub.font.color.rgb = ACCENT_PRIMARY
    p_sub.alignment = PP_ALIGN.CENTER

    # 4 Meta Information Cards
    meta_info = [
        ("TEAM MEMBERS", "Atharv Jain\nYashvardhan", ACCENT_CYAN),
        ("UNIVERSITY", "BML Munjal\nUniversity", ACCENT_PRIMARY),
        ("COURSE / SUBJECT", "AI DevOps\n(v3.0 Architecture)", ACCENT_PURPLE),
        ("FACULTY MENTOR", "Mr. Naveen Singh\nAcademic Advisor", ACCENT_EMERALD),
    ]

    card_w = Inches(2.55)
    gap = Inches(0.24)
    start_x = Inches(1.2)
    top_y = Inches(5.3)

    for i, (label, val, accent) in enumerate(meta_info):
        cx = start_x + i * (card_w + gap)
        mtf = create_card_with_text(s1, cx, top_y, card_w, Inches(1.5), bg_color=CARD_BG, border_color=CARD_BORDER, accent_bar=accent, pad_x=0.18, pad_y=0.14)
        
        mp0 = mtf.paragraphs[0]
        mp0.text = label
        mp0.font.name = FONT_BODY
        mp0.font.size = Pt(10)
        mp0.font.bold = True
        mp0.font.color.rgb = accent
        mp0.space_after = Pt(4)

        mp1 = mtf.add_paragraph()
        mp1.text = val
        mp1.font.name = FONT_TITLE
        mp1.font.size = Pt(12)
        mp1.font.bold = True
        mp1.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 2: Problem → Aim
    # =========================================================================
    s2 = create_blank_slide(prs)
    add_header(s2, "Problem → Aim: Why DevAssist AI?", "STRATEGIC CONTEXT & MISSION", 2)

    # Left: The Problem Card (Single text frame for all problems)
    ptf = create_card_with_text(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.4), accent_bar=ACCENT_ROSE, pad_x=0.25, pad_y=0.2)
    p_h = ptf.paragraphs[0]
    p_h.text = "The Core Problem"
    p_h.font.name = FONT_TITLE
    p_h.font.size = Pt(16)
    p_h.font.bold = True
    p_h.font.color.rgb = TEXT_WHITE
    p_h.space_after = Pt(2)

    p_sub = ptf.add_paragraph()
    p_sub.text = "Manual Code Reviews Cannot Scale With Rapid Agile Sprints"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(10.5)
    p_sub.font.color.rgb = ACCENT_ROSE
    p_sub.space_after = Pt(14)

    problems = [
        ("PR Review Bottlenecks", "Engineers spend 15-20% of engineering bandwidth manually reading pull requests, stalling delivery pipelines."),
        ("Multi-Language Blindspots", "Modern monorepos combine Python, TypeScript, and Go. Reviewers lack syntax familiarity across all."),
        ("Silent Downstream Blast Radius", "Modifying a function signature in a base module inadvertently breaks 15+ external callers across the repo."),
        ("Superficial Linter Fatigue", "Developers debate formatting and style comments rather than focusing on security, architecture, and race conditions."),
        ("Hallucinating First-Gen AI Bots", "Single-prompt LLM review bots flood developers with noisy, false-positive inline review comments."),
    ]
    for title, desc in problems:
        pt = ptf.add_paragraph()
        pt.text = f"❌  {title}"
        pt.font.name = FONT_TITLE
        pt.font.size = Pt(11)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_WHITE
        pt.space_after = Pt(1)

        pd = ptf.add_paragraph()
        pd.text = desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(7)

    # Right: The Aim Card (Top Box: Vision Statement)
    atf = create_card_with_text(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(2.2), accent_bar=ACCENT_EMERALD, pad_x=0.25, pad_y=0.18)
    a_h = atf.paragraphs[0]
    a_h.text = "The Engineering Aim"
    a_h.font.name = FONT_TITLE
    a_h.font.size = Pt(16)
    a_h.font.bold = True
    a_h.font.color.rgb = TEXT_WHITE
    a_h.space_after = Pt(2)

    a_sub = atf.add_paragraph()
    a_sub.text = "Autonomous, Evidence-Backed PR Review & Architecture Synthesis"
    a_sub.font.name = FONT_BODY
    a_sub.font.size = Pt(10.5)
    a_sub.font.color.rgb = ACCENT_EMERALD
    a_sub.space_after = Pt(8)

    a_body = atf.add_paragraph()
    a_body.text = "DevAssist AI aims to eliminate pull request bottlenecks by integrating a 3-stage multi-LLM ensemble pipeline, AST-aware CodeGraph blast radius calculation, and Hybrid RAG context to identify evidence-backed code defects, propose executable Git patch fixes, and synthesize architectural documentation in real-time."
    a_body.font.name = FONT_BODY
    a_body.font.size = Pt(9.5)
    a_body.font.color.rgb = TEXT_MUTED
    a_body.line_spacing = 1.2

    # Right Bottom: 4 Distinct Pillars (2x2 Grid)
    pillars = [
        ("⚡ FAST REVIEWS", "Sub-60s async review completion via Celery background tasks.", ACCENT_CYAN),
        ("🧠 MULTI-MODEL", "Specialized Distill-Reason-Validate stages prevent single-LLM errors.", ACCENT_PRIMARY),
        ("🔍 EVIDENCE-BACKED", "CodeGraph AST and linters verify every line citation and diff hunk.", ACCENT_PURPLE),
        ("👤 HUMAN-IN-CONTROL", "Developers review, accept, or dismiss findings; final merge stays human.", ACCENT_EMERALD),
    ]

    for idx, (p_title, p_desc, p_color) in enumerate(pillars):
        px = Inches(6.8) if idx % 2 == 0 else Inches(9.75)
        py = Inches(4.0) if idx < 2 else Inches(5.45)
        pw = Inches(2.75)
        ph = Inches(1.35)

        pctf = create_card_with_text(s2, px, py, pw, ph, bg_color=CARD_BG_HOVER, pad_x=0.18, pad_y=0.14)
        pp0 = pctf.paragraphs[0]
        pp0.text = p_title
        pp0.font.name = FONT_BODY
        pp0.font.size = Pt(10.5)
        pp0.font.bold = True
        pp0.font.color.rgb = p_color
        pp0.space_after = Pt(3)

        pp1 = pctf.add_paragraph()
        pp1.text = p_desc
        pp1.font.name = FONT_BODY
        pp1.font.size = Pt(9)
        pp1.font.color.rgb = TEXT_MUTED
        pp1.line_spacing = 1.15

    # =========================================================================
    # SLIDE 3: Multi-Agent & Multi-LLM Review System
    # =========================================================================
    s3 = create_blank_slide(prs)
    add_header(s3, "Our Idea: Multi-Agent & Multi-LLM Review Architecture", "MODULAR INTELLIGENCE SYSTEM", 3)

    agents = [
        ("1. Context Distiller", "Gemini 2.5 Flash", "Compresses diffs, linters, and blast radius into an executive risk digest.", ACCENT_CYAN),
        ("2. Deep Reasoner", "Claude 3.5 Sonnet", "Investigates race conditions, memory leaks, edge cases, and writes code fixes.", ACCENT_PRIMARY),
        ("3. Finding Validator", "Gemini 2.5 Flash", "Filters false positives; prunes findings scoring below 0.6 confidence.", ACCENT_EMERALD),
        ("4. CodeGraph Engine", "Tree-Sitter / AST", "Builds directed symbol call graph to detect affected downstream callers.", ACCENT_PURPLE),
        ("5. Hybrid RAG", "FAISS + BM25 + RRF", "Merges dense vector semantic retrieval with exact sparse keyword matching.", ACCENT_CYAN),
        ("6. Static Linters", "Pylint / ESLint", "Executes automated linters and passes rule violations directly to prompt.", ACCENT_AMBER),
        ("7. Agentic Sandbox", "Docker (Isolated)", "Generates and runs PR-specific Python verification scripts in a sandbox.", ACCENT_ROSE),
        ("8. Conversation Agent", "Multi-LLM Router", "Maintains thread context to answer developer replies directly on PR comments.", ACCENT_PRIMARY),
        ("9. Doc Synthesizer", "Doc Agent", "Autonomously authors Google/JSDoc docstrings and module documentation post-merge.", ACCENT_EMERALD),
    ]

    gw = Inches(3.7)
    gh = Inches(1.6)
    gx_start = Inches(0.8)
    gy_start = Inches(1.6)
    x_gap = Inches(0.3)
    y_gap = Inches(0.25)

    for i, (a_name, a_model, a_desc, a_acc) in enumerate(agents):
        row = i // 3
        col = i % 3
        x = gx_start + col * (gw + x_gap)
        y = gy_start + row * (gh + y_gap)

        actf = create_card_with_text(s3, x, y, gw, gh, accent_bar=a_acc, pad_x=0.2, pad_y=0.14)

        ap0 = actf.paragraphs[0]
        ap0.text = a_name
        ap0.font.name = FONT_TITLE
        ap0.font.size = Pt(12)
        ap0.font.bold = True
        ap0.font.color.rgb = TEXT_WHITE
        ap0.space_after = Pt(2)

        ap1 = actf.add_paragraph()
        ap1.text = f"Engine: {a_model}"
        ap1.font.name = FONT_MONO
        ap1.font.size = Pt(9.5)
        ap1.font.bold = True
        ap1.font.color.rgb = a_acc
        ap1.space_after = Pt(4)

        ap2 = actf.add_paragraph()
        ap2.text = a_desc
        ap2.font.name = FONT_BODY
        ap2.font.size = Pt(9)
        ap2.font.color.rgb = TEXT_MUTED
        ap2.line_spacing = 1.15

    # =========================================================================
    # SLIDE 4: How DevAssist AI Works – End-to-End
    # =========================================================================
    s4 = create_blank_slide(prs)
    add_header(s4, "How DevAssist AI Works – End-to-End Execution Flow", "SYSTEM PIPELINE & ORCHESTRATION", 4)

    steps = [
        ("1. PR Ingress", "GitHub / GitLab Webhook", "HMAC-SHA256 signature verification & debounce filter.", ACCENT_CYAN),
        ("2. Task Queue", "Redis + Celery", "Non-blocking dispatch; initiates background worker task.", ACCENT_PRIMARY),
        ("3. Code Analysis", "CodeGraph & Linters", "BFS caller traversal & parallel static code linting.", ACCENT_PURPLE),
        ("4. Context RAG", "AST Chunker + RRF", "Retrieves atomic function/class chunks from FAISS & BM25.", ACCENT_EMERALD),
        ("5. Distill Stage", "Gemini 2.5 Flash", "Compresses diffs & context into structured risk digest.", ACCENT_CYAN),
        ("6. Reason Stage", "Claude 3.5 Sonnet", "Detects bugs, security CVEs, and generates code fixes.", ACCENT_PRIMARY),
        ("7. Validate Stage", "Quality Gate (<0.6 drop)", "Hallucination filter scores confidence & verifies hunk lines.", ACCENT_AMBER),
        ("8. Output Delivery", "GitHub PR & Next.js", "Inline patch comments posted; live SSE stream to UI.", ACCENT_EMERALD),
    ]

    sw = Inches(2.75)
    sh = Inches(1.7)
    sx_start = Inches(0.8)
    sy_start = Inches(1.6)
    sgap = Inches(0.24)

    for i, (stitle, ssub, sbody, sacc) in enumerate(steps):
        row = i // 4
        col = i % 4
        x = sx_start + col * (sw + sgap)
        y = sy_start + row * (sh + Inches(0.25))

        ctf = create_card_with_text(s4, x, y, sw, sh, accent_bar=sacc, pad_x=0.18, pad_y=0.14)

        cp0 = ctf.paragraphs[0]
        cp0.text = stitle
        cp0.font.name = FONT_TITLE
        cp0.font.size = Pt(12)
        cp0.font.bold = True
        cp0.font.color.rgb = TEXT_WHITE
        cp0.space_after = Pt(2)

        cp1 = ctf.add_paragraph()
        cp1.text = ssub
        cp1.font.name = FONT_BODY
        cp1.font.size = Pt(9)
        cp1.font.bold = True
        cp1.font.color.rgb = sacc
        cp1.space_after = Pt(4)

        cp2 = ctf.add_paragraph()
        cp2.text = sbody
        cp2.font.name = FONT_BODY
        cp2.font.size = Pt(8.5)
        cp2.font.color.rgb = TEXT_MUTED
        cp2.line_spacing = 1.15

    # Bottom 3 Architecture Pillars
    bottom_cards = [
        ("AST Chunking & CodeGraph", "Preserves function and class boundaries instead of arbitrary character slices. CodeGraph tracks callers up to 3 hops deep.", ACCENT_CYAN),
        ("Asynchronous Celery Workers", "Background Celery workers process files concurrently using thread pools, keeping the FastAPI webhook ingress fully non-blocking.", ACCENT_PRIMARY),
        ("Cross-Process SSE Telemetry", "Celery worker processes publish events to Redis Pub/Sub channels (`devassist:sse:events`), streaming real-time status to the browser.", ACCENT_EMERALD),
    ]

    bw = Inches(3.7)
    bh = Inches(1.4)
    by = Inches(5.55)

    for i, (btitle, bdesc, bacc) in enumerate(bottom_cards):
        bx = sx_start + i * (bw + Inches(0.31))
        bctf = create_card_with_text(s4, bx, by, bw, bh, bg_color=CARD_BG_HOVER, pad_x=0.18, pad_y=0.14)

        bp0 = bctf.paragraphs[0]
        bp0.text = btitle
        bp0.font.name = FONT_TITLE
        bp0.font.size = Pt(11)
        bp0.font.bold = True
        bp0.font.color.rgb = bacc
        bp0.space_after = Pt(3)

        bp1 = bctf.add_paragraph()
        bp1.text = bdesc
        bp1.font.name = FONT_BODY
        bp1.font.size = Pt(9)
        bp1.font.color.rgb = TEXT_MUTED
        bp1.line_spacing = 1.15

    # =========================================================================
    # SLIDE 5: From Findings to Decisions
    # =========================================================================
    s5 = create_blank_slide(prs)
    add_header(s5, "From Findings to Decisions: Triage & Quality Gate", "EVALUATION & VERIFICATION", 5)

    col_w = Inches(3.7)
    col_h = Inches(5.4)
    cy = Inches(1.6)

    # Col 1: Blast Radius
    c1_tf = create_card_with_text(s5, Inches(0.8), cy, col_w, col_h, accent_bar=ACCENT_ROSE, pad_x=0.22, pad_y=0.18)
    p0 = c1_tf.paragraphs[0]
    p0.text = "1. CodeGraph Blast Radius"
    p0.font.name = FONT_TITLE
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_WHITE
    p0.space_after = Pt(2)

    p_sub = c1_tf.add_paragraph()
    p_sub.text = "Downstream Caller Impact Analysis"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(10)
    p_sub.font.color.rgb = ACCENT_ROSE
    p_sub.space_after = Pt(12)

    c1_pts = [
        ("Changed Symbol", "Function `verify_token()` modified in `auth/jwt.py`"),
        ("BFS Traversal", "Graph explorer walks callers recursively up to 3 hops deep."),
        ("Transitive Callers", "14 direct and indirect caller endpoints identified across API routes."),
        ("Severity Assessment", "Label: HIGH BLAST RADIUS (>6 callers impacted)."),
        ("Developer Value", "Prevents breaking public API contracts before merging to staging."),
    ]
    for h, b in c1_pts:
        p = c1_tf.add_paragraph()
        p.text = f"• {h}"
        p.font.name = FONT_TITLE
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(1)

        pd = c1_tf.add_paragraph()
        pd.text = b
        pd.font.name = FONT_BODY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(6)

    # Col 2: Quality Gate
    c2_tf = create_card_with_text(s5, Inches(4.8), cy, col_w, col_h, accent_bar=ACCENT_AMBER, pad_x=0.22, pad_y=0.18)
    p0 = c2_tf.paragraphs[0]
    p0.text = "2. Quality Gate (Validator)"
    p0.font.name = FONT_TITLE
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_WHITE
    p0.space_after = Pt(2)

    p_sub = c2_tf.add_paragraph()
    p_sub.text = "Confidence & Hallucination Filter"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(10)
    p_sub.font.color.rgb = ACCENT_AMBER
    p_sub.space_after = Pt(12)

    c2_pts = [
        ("Hunk Line Verification", "Ensures finding target line exists in the PR diff hunk."),
        ("Executable Fix Check", "Validates syntax of suggested code replacement."),
        ("Confidence Scoring", "Validator scores finding validity between 0.0 and 1.0."),
        ("Pruning Rule", "Findings below 0.6 threshold are automatically dropped."),
        ("Deduplication", "Suppresses duplicate findings across consecutive commits."),
    ]
    for h, b in c2_pts:
        p = c2_tf.add_paragraph()
        p.text = f"✔ {h}"
        p.font.name = FONT_TITLE
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = ACCENT_EMERALD
        p.space_after = Pt(1)

        pd = c2_tf.add_paragraph()
        pd.text = b
        pd.font.name = FONT_BODY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(6)

    # Col 3: PR Quality Scoring
    c3_tf = create_card_with_text(s5, Inches(8.8), cy, col_w, col_h, accent_bar=ACCENT_EMERALD, pad_x=0.22, pad_y=0.18)
    p0 = c3_tf.paragraphs[0]
    p0.text = "3. PR Quality Scoring"
    p0.font.name = FONT_TITLE
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_WHITE
    p0.space_after = Pt(2)

    p_sub = c3_tf.add_paragraph()
    p_sub.text = "Deterministic Metric & Triage"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(10)
    p_sub.font.color.rgb = ACCENT_EMERALD
    p_sub.space_after = Pt(12)

    c3_pts = [
        ("🔴 Errors (-2.0 pts)", "Security vulnerabilities, data corruption, uncaught exceptions."),
        ("🟡 Warnings (-0.5 pts)", "Performance regressions, missing timeouts, edge cases."),
        ("🔵 Notes (-0.1 pts)", "Documentation gaps, style inconsistencies, cleanups."),
        ("Score Formula", "Quality Score = max(0, 10 - Penalty Score). Range 0 to 10."),
        ("Engineer in Control", "Engineers can suppress false positives or approve PR."),
    ]
    for h, b in c3_pts:
        p = c3_tf.add_paragraph()
        p.text = h
        p.font.name = FONT_TITLE
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(1)

        pd = c3_tf.add_paragraph()
        pd.text = b
        pd.font.name = FONT_BODY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(6)

    # =========================================================================
    # SLIDE 6: System Architecture & Technology Stack
    # =========================================================================
    s6 = create_blank_slide(prs)
    add_header(s6, "System Architecture & Comprehensive Tech Stack", "FULL-STACK CLOUD ARCHITECTURE", 6)

    tech_pillars = [
        ("Frontend & Dashboard", "Next.js 16 + React 19", [
            ("Framework", "Next.js 16 App Router (Turbopack)"),
            ("UI Library", "React 19, TypeScript 5, Tailwind v4"),
            ("Diff Viewer", "@pierre/diffs side-by-side viewer"),
            ("Visual Graph", "ReactFlow interactive blast radius"),
            ("Diagrams", "Mermaid live sequence & flow diagrams"),
            ("Auth", "Clerk enterprise user management"),
        ], ACCENT_CYAN),

        ("Backend & Distributed Queue", "FastAPI + Celery + Redis", [
            ("API Server", "FastAPI v3.0 (async routes)"),
            ("Task Worker", "Celery 5.3+ solo/fork worker pool"),
            ("Broker", "Redis 7 (Queue + Pub/Sub SSE)"),
            ("Database", "PostgreSQL 15 (asyncpg) / SQLite"),
            ("ORM Layer", "SQLAlchemy 2.0 async sessions"),
            ("Migrations", "Alembic automated schema tracking"),
        ], ACCENT_PRIMARY),

        ("Multi-LLM & Context Engines", "Multi-Stage Intelligence", [
            ("Reasoning LLM", "Claude 3.5 Sonnet / GPT-4o"),
            ("Fast Distiller", "Google Gemini 2.5 Flash"),
            ("Fallback Router", "NVIDIA NIM, OpenRouter, Local Ollama"),
            ("Dense Index", "FAISS CPU + Sentence Transformers"),
            ("Sparse Index", "rank-bm25 keyword search index"),
            ("Fusion", "Reciprocal Rank Fusion (RRF)"),
        ], ACCENT_PURPLE),

        ("DevOps, Sandbox & Ingress", "Containerized Infrastructure", [
            ("Containers", "Docker & Docker Compose stack"),
            ("Sandbox", "Isolated Docker (--network none)"),
            ("Security", "nobody user, dropped Linux caps"),
            ("Ingress", "GitHub App + HMAC-SHA256 signature"),
            ("Tunneling", "Pyngrok automated public tunnel"),
            ("Linters", "Parallel Pylint, ESLint, Checkstyle"),
        ], ACCENT_EMERALD),
    ]

    pw = Inches(2.75)
    ph = Inches(5.4)
    px_start = Inches(0.8)
    py_start = Inches(1.6)
    pgap = Inches(0.24)

    for i, (t_title, t_sub, items, t_acc) in enumerate(tech_pillars):
        x = px_start + i * (pw + pgap)
        itf = create_card_with_text(s6, x, py_start, pw, ph, accent_bar=t_acc, pad_x=0.18, pad_y=0.16)

        p0 = itf.paragraphs[0]
        p0.text = t_title
        p0.font.name = FONT_TITLE
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_WHITE
        p0.space_after = Pt(2)

        p1 = itf.add_paragraph()
        p1.text = t_sub
        p1.font.name = FONT_BODY
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = t_acc
        p1.space_after = Pt(10)

        for k, v in items:
            pk = itf.add_paragraph()
            pk.text = k
            pk.font.name = FONT_TITLE
            pk.font.size = Pt(9.5)
            pk.font.bold = True
            pk.font.color.rgb = t_acc
            pk.space_after = Pt(1)

            pv = itf.add_paragraph()
            pv.text = v
            pv.font.name = FONT_BODY
            pv.font.size = Pt(8.5)
            pv.font.color.rgb = TEXT_WHITE
            pv.space_after = Pt(5)

    # =========================================================================
    # SLIDE 7: What the User Sees
    # =========================================================================
    s7 = create_blank_slide(prs)
    add_header(s7, "What the User Sees: Developer Dashboard & Native PR UX", "USER EXPERIENCE & WORKFLOW", 7)

    # Left: Web Dashboard
    dtf = create_card_with_text(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.4), accent_bar=ACCENT_CYAN, pad_x=0.25, pad_y=0.2)
    dh = dtf.paragraphs[0]
    dh.text = "Next.js 16 Web Dashboard"
    dh.font.name = FONT_TITLE
    dh.font.size = Pt(16)
    dh.font.bold = True
    dh.font.color.rgb = TEXT_WHITE
    dh.space_after = Pt(2)

    dsub = dtf.add_paragraph()
    dsub.text = "Centralized Command Center for Monitored Repositories"
    dsub.font.name = FONT_BODY
    dsub.font.size = Pt(10.5)
    dsub.font.color.rgb = ACCENT_CYAN
    dsub.space_after = Pt(14)

    dash_features = [
        ("Engineering Overview", "Real-time velocity tracking, review counts, and 30-day severity trends."),
        ("Interactive Blast Radius Graph", "ReactFlow dynamic node canvas displaying changed files and affected callers."),
        ("Pierre Diff Viewer", "Side-by-side syntax-highlighted diffs split into manageable chapter walkthroughs."),
        ("Architectural Prologue", "Live-rendered Mermaid sequence and flow diagrams detailing structural modifications."),
        ("MCP Tooling & Integrations", "Standardized Model Context Protocol integrations with Jira, Linear, and Notion."),
    ]
    for f_title, f_desc in dash_features:
        p = dtf.add_paragraph()
        p.text = f"📊  {f_title}"
        p.font.name = FONT_TITLE
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        p.space_after = Pt(1)

        pd = dtf.add_paragraph()
        pd.text = f_desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(8)

    # Right: Native GitHub PR Experience
    ghtf = create_card_with_text(s7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.4), accent_bar=ACCENT_PRIMARY, pad_x=0.25, pad_y=0.2)
    gh_h = ghtf.paragraphs[0]
    gh_h.text = "GitHub & GitLab PR Experience"
    gh_h.font.name = FONT_TITLE
    gh_h.font.size = Pt(16)
    gh_h.font.bold = True
    gh_h.font.color.rgb = TEXT_WHITE
    gh_h.space_after = Pt(2)

    gh_sub = ghtf.add_paragraph()
    gh_sub.text = "Frictionless CodeRabbit-Style Inline Reviews & Patches"
    gh_sub.font.name = FONT_BODY
    gh_sub.font.size = Pt(10.5)
    gh_sub.font.color.rgb = ACCENT_PRIMARY
    gh_sub.space_after = Pt(14)

    gh_features = [
        ("Inline Code Review Comments", "Precise line-level comments explaining bugs, security concerns, or performance pitfalls."),
        ("Suggested Git Patches (```suggestion)", "Developers can apply AI-authored fixes directly with a single click inside GitHub."),
        ("Summary Walkthrough Comment", "Detailed Markdown table summarizing changed files, additions/deletions, and blast radius."),
        ("PR Quality Score Banner", "Emoji-coded score badge (e.g. 🟢 8.5/10) with categorical finding breakdown."),
        ("Conversational Comment Replies", "Developers can reply directly to bot comments to ask for clarification or alternative fixes."),
    ]
    for g_title, g_desc in gh_features:
        p = ghtf.add_paragraph()
        p.text = f"🤖  {g_title}"
        p.font.name = FONT_TITLE
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_PRIMARY
        p.space_after = Pt(1)

        pd = ghtf.add_paragraph()
        pd.text = g_desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(8)

    # =========================================================================
    # SLIDE 8: Where AI Reduces Effort
    # =========================================================================
    s8 = create_blank_slide(prs)
    add_header(s8, "Where AI Reduces Effort: Manual vs DevAssist AI", "WORKFLOW TRANSFORMATION MATRIX", 8)

    # Left & Right Headers
    h_man_tf = create_card_with_text(s8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(0.52), bg_color=RGBColor(36, 18, 26), border_color=ACCENT_ROSE, pad_x=0.15, pad_y=0.12)
    mp = h_man_tf.paragraphs[0]
    mp.text = "❌  Traditional Manual Code Review Workflow"
    mp.font.name = FONT_TITLE
    mp.font.size = Pt(11.5)
    mp.font.bold = True
    mp.font.color.rgb = ACCENT_ROSE

    h_ai_tf = create_card_with_text(s8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(0.52), bg_color=RGBColor(16, 36, 30), border_color=ACCENT_EMERALD, pad_x=0.15, pad_y=0.12)
    ap = h_ai_tf.paragraphs[0]
    ap.text = "⚡  DevAssist AI Autonomous Assisted Workflow"
    ap.font.name = FONT_TITLE
    ap.font.size = Pt(11.5)
    ap.font.bold = True
    ap.font.color.rgb = ACCENT_EMERALD

    comparisons = [
        ("Manual Line-by-Line Reading", "Engineers read raw git diffs without codebase context, missing cross-file variable references.",
         "AST & Hybrid RAG Retrieval", "Atomic syntax chunking retrieves relevant classes, types, and module context automatically."),

        ("Missed Caller Impacts", "Transitive breaking changes in callers remain undiscovered until staging/production failures.",
         "CodeGraph BFS Blast Radius", "Traverses symbol dependency edges up to 3 hops deep and classifies caller impact severity."),

        ("Syntax & Linter Debates", "Reviewers spend time pointing out PEP8/ESLint issues that could be resolved automatically.",
         "Parallel Static & Dynamic Linters", "Runs Pylint, ESLint, and sandboxed Docker tests; injects rule IDs straight into prompts."),

        ("Unfiltered Bot Hallucinations", "Generic LLM bots hallucinate non-existent imports or suggest invalid line numbers.",
         "Stage 3 Finding Validator", "Checks diff hunk bounds, validates suggested code fix syntax, and prunes <0.6 confidence."),

        ("Neglected Documentation", "Docstrings and READMEs go stale as engineers rush to ship feature code.",
         "Autonomous Docstring Synthesizer", "Generates compliant docstrings and Markdown docs immediately upon PR merge."),
    ]

    row_y = Inches(2.25)
    row_h = Inches(0.85)
    rgap = Inches(0.12)

    for i, (m_title, m_desc, a_title, a_desc) in enumerate(comparisons):
        y = row_y + i * (row_h + rgap)

        mctf = create_card_with_text(s8, Inches(0.8), y, Inches(5.6), row_h, pad_x=0.16, pad_y=0.1)
        p0 = mctf.paragraphs[0]
        p0.text = m_title
        p0.font.name = FONT_TITLE
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_WHITE
        p0.space_after = Pt(1)

        p1 = mctf.add_paragraph()
        p1.text = m_desc
        p1.font.name = FONT_BODY
        p1.font.size = Pt(8.5)
        p1.font.color.rgb = TEXT_MUTED

        actf = create_card_with_text(s8, Inches(6.8), y, Inches(5.7), row_h, pad_x=0.16, pad_y=0.1)
        ap0 = actf.paragraphs[0]
        ap0.text = a_title
        ap0.font.name = FONT_TITLE
        ap0.font.size = Pt(10)
        ap0.font.bold = True
        ap0.font.color.rgb = ACCENT_EMERALD
        ap0.space_after = Pt(1)

        ap1 = actf.add_paragraph()
        ap1.text = a_desc
        ap1.font.name = FONT_BODY
        ap1.font.size = Pt(8.5)
        ap1.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 9: What Improves with DevAssist AI? (Metrics)
    # =========================================================================
    s9 = create_blank_slide(prs)
    add_header(s9, "What Improves with DevAssist AI? – Quantitative Impact", "BENCHMARKS & CORE VALUE DRIVERS", 9)

    metrics = [
        ("70%+", "FASTER PR TURNAROUND", "Celery async worker execution and concurrent thread-pool review reduce review waiting time from hours to under 60 seconds.", ACCENT_CYAN),
        ("< 5%", "FALSE POSITIVE RATE", "3-stage distillation and confidence scoring validator systematically prune ungrounded and noisy hallucinations.", ACCENT_EMERALD),
        ("3 HOPS", "BLAST RADIUS DEPTH", "AST-aware CodeGraph traverses transitive downstream callers up to 3 hops deep to catch breaking changes before merge.", ACCENT_PRIMARY),
        ("99.9%", "ROUTING AVAILABILITY", "Multi-LLM fallback chain automatically routes across OpenAI, Claude, Gemini, NVIDIA, and Ollama on error.", ACCENT_PURPLE),
        ("1-CLICK", "SUGGESTED PATCH FIXES", "Generates executable GitHub suggested code blocks (```suggestion) so developers can apply fixes instantly.", ACCENT_AMBER),
    ]

    mw = Inches(2.18)
    mh = Inches(5.2)
    mx_start = Inches(0.8)
    my = Inches(1.7)
    mgap = Inches(0.2)

    for i, (stat, stat_lbl, stat_desc, stat_acc) in enumerate(metrics):
        x = mx_start + i * (mw + mgap)
        ctf = create_card_with_text(s9, x, my, mw, mh, accent_bar=stat_acc, pad_x=0.16, pad_y=0.25)

        p0 = ctf.paragraphs[0]
        p0.text = stat
        p0.font.name = FONT_TITLE
        p0.font.size = Pt(32)
        p0.font.bold = True
        p0.font.color.rgb = stat_acc
        p0.alignment = PP_ALIGN.CENTER
        p0.space_after = Pt(4)

        p1 = ctf.add_paragraph()
        p1.text = stat_lbl
        p1.font.name = FONT_BODY
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.alignment = PP_ALIGN.CENTER
        p1.space_after = Pt(16)

        p2 = ctf.add_paragraph()
        p2.text = stat_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_MUTED
        p2.line_spacing = 1.25

    # =========================================================================
    # SLIDE 10: DevOps & Deployment
    # =========================================================================
    s10 = create_blank_slide(prs)
    add_header(s10, "DevOps & Deployment: From Development to Monitoring", "ENTERPRISE DEPLOYMENT & CI/CD", 10)

    devops_steps = [
        ("1. Git Commit", "Push code branch.", ACCENT_CYAN),
        ("2. Containerize", "Docker API & worker.", ACCENT_PRIMARY),
        ("3. Compose", "Redis, Postgres, Celery.", ACCENT_PURPLE),
        ("4. Tunneling", "Pyngrok webhook ingress.", ACCENT_AMBER),
        ("5. Sandbox", "Docker test runner.", ACCENT_ROSE),
        ("6. Telemetry", "Redis Pub/Sub SSE.", ACCENT_EMERALD),
    ]

    dw = Inches(1.8)
    dh = Inches(1.05)
    dx_start = Inches(0.8)
    dy = Inches(1.6)
    dgap = Inches(0.18)

    for i, (d_title, d_sub, d_acc) in enumerate(devops_steps):
        x = dx_start + i * (dw + dgap)
        ctf = create_card_with_text(s10, x, dy, dw, dh, accent_bar=d_acc, pad_x=0.12, pad_y=0.1)

        cp0 = ctf.paragraphs[0]
        cp0.text = d_title
        cp0.font.name = FONT_TITLE
        cp0.font.size = Pt(10)
        cp0.font.bold = True
        cp0.font.color.rgb = TEXT_WHITE
        cp0.space_after = Pt(2)

        cp1 = ctf.add_paragraph()
        cp1.text = d_sub
        cp1.font.name = FONT_BODY
        cp1.font.size = Pt(8.5)
        cp1.font.color.rgb = TEXT_MUTED

    devops_cards = [
        ("Docker Compose Production Stack", "Complete Multi-Service Orchestration", [
            ("Redis 7 Alpine", "In-memory broker for Celery tasks and cross-process SSE event propagation."),
            ("PostgreSQL 15", "Relational persistence storing repositories, PR metadata, reviews, and findings."),
            ("Celery Background Worker", "Dedicated task runner executing full PR review pipelines and RAG indexing."),
            ("FastAPI Backend", "Async REST server with live CORS, HMAC webhook validation, and health checks."),
        ], ACCENT_CYAN),

        ("Secure Agentic Execution Sandbox", "Zero-Trust Dynamic Script Execution", [
            ("Isolated Container", "Runs dynamic Python scripts inside a hardened container with --network none."),
            ("Unprivileged User", "Executes as nobody with all Linux capabilities dropped to prevent host escapes."),
            ("Read-Only Mounting", "The cloned repository is mounted read-only at /repo to avoid malicious mutations."),
            ("Wall-Clock Timeout", "Kills and prunes containers automatically if a test script exceeds 30 seconds."),
        ], ACCENT_ROSE),

        ("Observability & Telemetry", "Real-Time Health & Audit Logging", [
            ("Live /health Endpoint", "Monitors active database connection, active LLM provider, and queue connectivity."),
            ("Redis Pub/Sub SSE", "Streams stage transitions (`distill`, `reason`, `validate`) directly to Next.js UI."),
            ("Alembic Migrations", "Automated database schema versioning for zero-downtime updates."),
            ("Audit Event Trail", "Persistent event logs record all review actions, findings, and developer dismissals."),
        ], ACCENT_EMERALD),
    ]

    cw = Inches(3.7)
    ch = Inches(4.2)
    cy = Inches(2.85)

    for i, (title, sub, bullet_items, acc) in enumerate(devops_cards):
        cx = Inches(0.8) + i * (cw + Inches(0.31))
        ctf = create_card_with_text(s10, cx, cy, cw, ch, accent_bar=acc, pad_x=0.2, pad_y=0.18)

        p0 = ctf.paragraphs[0]
        p0.text = title
        p0.font.name = FONT_TITLE
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_WHITE
        p0.space_after = Pt(2)

        p1 = ctf.add_paragraph()
        p1.text = sub
        p1.font.name = FONT_BODY
        p1.font.size = Pt(9.5)
        p1.font.color.rgb = acc
        p1.space_after = Pt(10)

        for k, v in bullet_items:
            pk = ctf.add_paragraph()
            pk.text = f"• {k}"
            pk.font.name = FONT_TITLE
            pk.font.size = Pt(9.5)
            pk.font.bold = True
            pk.font.color.rgb = acc
            pk.space_after = Pt(1)

            pv = ctf.add_paragraph()
            pv.text = v
            pv.font.name = FONT_BODY
            pv.font.size = Pt(8.5)
            pv.font.color.rgb = TEXT_MUTED
            pv.space_after = Pt(4)

    # =========================================================================
    # SLIDE 11: Future Scope
    # =========================================================================
    s11 = create_blank_slide(prs)
    add_header(s11, "Future Scope: Roadmap & Next-Gen Capabilities", "STRATEGIC HORIZON & EXPANSION", 11)

    future_items = [
        ("IDE Live Pre-Commit Plugins", "VS Code & JetBrains Extensions", "Run lightweight fast-mode reviews locally in the IDE before commits are pushed to remote branches.", ACCENT_CYAN),
        ("Autonomous Self-Healing PRs", "Auto-Fix Commit Branching", "Automatically create verified fix branches with passing unit test suites for straightforward bug fixes.", ACCENT_PRIMARY),
        ("Cross-Repo Monorepo Graphs", "Microservice Contract Tracing", "Expand CodeGraph to index gRPC and OpenAPI contracts across distributed multi-repository architectures.", ACCENT_PURPLE),
        ("Fine-Tuned Enterprise LLMs", "Private On-Premise Models", "Train domain-specific LoRA adapters on internal enterprise coding guidelines and security playbooks.", ACCENT_EMERALD),
        ("Enterprise RBAC & Governance", "SOC2 Compliance & SAML SSO", "Role-based access control, exportable compliance audit logs, and custom regulatory policy enforcement.", ACCENT_AMBER),
        ("Automated Regression Test Gen", "Diff-Targeted Test Creation", "Synthesize unit and integration tests covering newly introduced diff logic to prevent regressions.", ACCENT_ROSE),
    ]

    fw = Inches(3.7)
    fh = Inches(2.25)
    fx_start = Inches(0.8)
    fy_start = Inches(1.6)
    fgap_x = Inches(0.31)
    fgap_y = Inches(0.25)

    for i, (f_title, f_sub, f_desc, f_acc) in enumerate(future_items):
        row = i // 3
        col = i % 3
        x = fx_start + col * (fw + fgap_x)
        y = fy_start + row * (fh + fgap_y)

        ctf = create_card_with_text(s11, x, y, fw, fh, accent_bar=f_acc, pad_x=0.2, pad_y=0.14)

        p0 = ctf.paragraphs[0]
        p0.text = f_title
        p0.font.name = FONT_TITLE
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_WHITE
        p0.space_after = Pt(2)

        p1 = ctf.add_paragraph()
        p1.text = f_sub
        p1.font.name = FONT_BODY
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = f_acc
        p1.space_after = Pt(6)

        p2 = ctf.add_paragraph()
        p2.text = f_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9)
        p2.font.color.rgb = TEXT_MUTED
        p2.line_spacing = 1.15

    # Bottom Goal Banner
    mbtf = create_card_with_text(s11, Inches(0.8), Inches(6.45), Inches(11.733), Inches(0.65), bg_color=CARD_BG_HOVER, border_color=ACCENT_PRIMARY, pad_x=0.2, pad_y=0.12)
    mp = mbtf.paragraphs[0]
    mp.text = "🎯  Strategic Goal: Make DevAssist AI the most context-aware, reliable, and secure autonomous engineering assistant for enterprise software teams."
    mp.font.name = FONT_BODY
    mp.font.size = Pt(10)
    mp.font.bold = True
    mp.font.color.rgb = TEXT_WHITE
    mp.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 12: Thank You & Q&A
    # =========================================================================
    s12 = create_blank_slide(prs)

    # Hero Box
    hero12 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.0), Inches(10.333), Inches(5.4))
    hero12.fill.solid()
    hero12.fill.fore_color.rgb = CARD_BG
    hero12.line.color.rgb = ACCENT_PRIMARY
    hero12.line.width = Pt(1.5)

    ostrip = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.0), Inches(10.333), Inches(0.08))
    ostrip.fill.solid()
    ostrip.fill.fore_color.rgb = ACCENT_CYAN
    ostrip.line.fill.background()

    # Thank you texts in single text frame
    ot_box = s12.shapes.add_textbox(Inches(2.0), Inches(1.4), Inches(9.333), Inches(2.4))
    otf = ot_box.text_frame
    otf.word_wrap = True
    otf.margin_left = 0
    otf.margin_top = 0
    otf.margin_right = 0
    otf.margin_bottom = 0

    op0 = otf.paragraphs[0]
    op0.text = "Thank You!"
    op0.font.name = FONT_TITLE
    op0.font.size = Pt(44)
    op0.font.bold = True
    op0.font.color.rgb = TEXT_WHITE
    op0.alignment = PP_ALIGN.CENTER
    op0.space_after = Pt(8)

    op1 = otf.add_paragraph()
    op1.text = "DevAssist AI – Autonomous Code Review & Architectural Documentation"
    op1.font.name = FONT_BODY
    op1.font.size = Pt(15)
    op1.font.bold = True
    op1.font.color.rgb = ACCENT_PRIMARY
    op1.alignment = PP_ALIGN.CENTER
    op1.space_after = Pt(4)

    op2 = otf.add_paragraph()
    op2.text = "Open for Questions, Feedback & Live Demonstration"
    op2.font.name = FONT_BODY
    op2.font.size = Pt(12)
    op2.font.color.rgb = ACCENT_CYAN
    op2.alignment = PP_ALIGN.CENTER

    # 3 Credit Cards
    credits = [
        ("PROJECT AUTHORS", "Atharv Jain\nYashvardhan", "AI DevOps Engineering Team", ACCENT_CYAN),
        ("INSTITUTION", "BML Munjal University", "Department of Computer Science", ACCENT_PRIMARY),
        ("FACULTY MENTOR", "Mr. Naveen Singh", "Project Advisor & Guide", ACCENT_EMERALD),
    ]

    cw = Inches(2.8)
    cy = Inches(4.3)
    cgap = Inches(0.4)
    cx_start = Inches(2.06)

    for i, (cr_role, cr_name, cr_sub, cr_acc) in enumerate(credits):
        x = cx_start + i * (cw + cgap)
        cctf = create_card_with_text(s12, x, cy, cw, Inches(1.6), bg_color=CARD_BG_HOVER, accent_bar=cr_acc, pad_x=0.15, pad_y=0.16)

        cp0 = cctf.paragraphs[0]
        cp0.text = cr_role
        cp0.font.name = FONT_BODY
        cp0.font.size = Pt(9.5)
        cp0.font.bold = True
        cp0.font.color.rgb = cr_acc
        cp0.alignment = PP_ALIGN.CENTER
        cp0.space_after = Pt(4)

        cp1 = cctf.add_paragraph()
        cp1.text = cr_name
        cp1.font.name = FONT_TITLE
        cp1.font.size = Pt(12)
        cp1.font.bold = True
        cp1.font.color.rgb = TEXT_WHITE
        cp1.alignment = PP_ALIGN.CENTER
        cp1.space_after = Pt(2)

        cp2 = cctf.add_paragraph()
        cp2.text = cr_sub
        cp2.font.name = FONT_BODY
        cp2.font.size = Pt(8.5)
        cp2.font.color.rgb = TEXT_MUTED
        cp2.alignment = PP_ALIGN.CENTER

    prs.save(out_path)
    print(f"Clean presentation successfully built and saved to: {out_path}")

if __name__ == "__main__":
    out_file = "E:/Devassist-AI-main/devassist-ai/DevAssist_AI_Clean_Presentation.pptx"
    build_presentation(out_file)
