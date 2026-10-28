import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Color Palette (Modern Dark / Cyber Indigo Theme) ──────────────────────────
BG_COLOR         = RGBColor(10, 13, 22)      # #0A0D16 - Deep midnight obsidian
CARD_BG          = RGBColor(18, 24, 40)      # #121828 - Elevated dark card
CARD_BG_HOVER    = RGBColor(24, 32, 54)      # #182036 - Lighter surface
CARD_BORDER      = RGBColor(38, 52, 84)      # #263454 - Subtle card border
ACCENT_PRIMARY   = RGBColor(99, 102, 241)    # #6366F1 - Electric Indigo
ACCENT_CYAN      = RGBColor(14, 165, 233)    # #0EA5E9 - Cyber Blue
ACCENT_EMERALD   = RGBColor(16, 185, 129)    # #10B981 - Success Green
ACCENT_ROSE      = RGBColor(244, 63, 94)     # #F43F5E - Error / Alert Rose
ACCENT_AMBER     = RGBColor(245, 158, 11)    # #F59E0B - Warning Amber
ACCENT_PURPLE    = RGBColor(168, 85, 247)    # #A855F7 - Modern Violet

TEXT_WHITE       = RGBColor(248, 250, 252)   # #F8FAFC - Pure bright text
TEXT_MUTED       = RGBColor(148, 163, 184)   # #94A3B8 - Slate body text
TEXT_DIM         = RGBColor(100, 116, 139)   # #64748B - Subtitle / footnote text
BADGE_BG         = RGBColor(30, 41, 69)      # #1E2945 - Pill badge fill

FONT_TITLE = "Segoe UI"
FONT_BODY  = "Segoe UI"
FONT_MONO  = "Consolas"

def create_blank_slide(prs):
    """Creates a blank slide with the custom deep midnight background."""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()
    return slide

def add_header(slide, title_text, category_badge="DEVASSIST AI PLATFORM", slide_number=None):
    """Standardized modern header with glowing category badge, title, and bottom divider."""
    # Category badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(2.9), Inches(0.32))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BADGE_BG
    badge.line.color.rgb = ACCENT_PRIMARY
    badge.line.width = Pt(1)
    tf = badge.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.1)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.04)
    p = tf.paragraphs[0]
    p.text = category_badge
    p.font.name = FONT_BODY
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.alignment = PP_ALIGN.CENTER

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.0), Inches(0.6))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = 0
    tf_t.margin_top = 0
    tf_t.margin_right = 0
    tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = FONT_TITLE
    p_t.font.size = Pt(23)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_WHITE

    # Slide Number pill (top right)
    if slide_number:
        num_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.933), Inches(0.35), Inches(0.6), Inches(0.32))
        num_box.fill.solid()
        num_box.fill.fore_color.rgb = BADGE_BG
        num_box.line.color.rgb = CARD_BORDER
        tf_n = num_box.text_frame
        tf_n.margin_top = Inches(0.04)
        pn = tf_n.paragraphs[0]
        pn.text = str(slide_number).zfill(2)
        pn.font.name = FONT_MONO
        pn.font.size = Pt(10)
        pn.font.bold = True
        pn.font.color.rgb = TEXT_MUTED
        pn.alignment = PP_ALIGN.CENTER

    # Thin glowing accent divider line
    div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.38), Inches(11.733), Inches(0.02))
    div.fill.solid()
    div.fill.fore_color.rgb = CARD_BORDER
    div.line.fill.background()

def create_card_with_text(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, accent_bar=None, pad_x=0.22, pad_y=0.18):
    """
    Creates an empty card container and returns a clean, single text frame on top.
    By putting all text in this SINGLE text frame, text elements can NEVER overlap or collide.
    """
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)

    top_offset = pad_y
    if accent_bar:
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent_bar
        bar.line.fill.background()
        top_offset += 0.06

    tb = slide.shapes.add_textbox(
        left + Inches(pad_x),
        top + Inches(top_offset),
        width - Inches(pad_x * 2),
        height - Inches(top_offset + pad_y)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_top = 0
    tf.margin_right = 0
    tf.margin_bottom = 0
    return tf

# ── Main Presentation Builder ─────────────────────────────────────────────────

def build_presentation(out_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # =========================================================================
    # SLIDE 1: Title Slide (Hero)
    # =========================================================================
    s1 = create_blank_slide(prs)

    # Hero Background Box
    hero = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(0.8), Inches(10.933), Inches(4.2))
    hero.fill.solid()
    hero.fill.fore_color.rgb = CARD_BG
    hero.line.color.rgb = ACCENT_PRIMARY
    hero.line.width = Pt(1.5)

    # Top Cyan Strip
    strip = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(0.8), Inches(10.933), Inches(0.08))
    strip.fill.solid()
    strip.fill.fore_color.rgb = ACCENT_CYAN
    strip.line.fill.background()

    # Single text frame for hero content to eliminate any vertical collisions
    h_tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.2), Inches(10.333), Inches(3.6))
    h_tf = h_tb.text_frame
    h_tf.word_wrap = True
    h_tf.margin_left = 0
    h_tf.margin_top = 0
    h_tf.margin_right = 0
    h_tf.margin_bottom = 0

    p_pill = h_tf.paragraphs[0]
    p_pill.text = "[ SMART AI ASSISTANT FOR AUTOMATIC CODE REVIEW ]"
    p_pill.font.name = FONT_BODY
    p_pill.font.size = Pt(12)
    p_pill.font.bold = True
    p_pill.font.color.rgb = ACCENT_CYAN
    p_pill.alignment = PP_ALIGN.CENTER
    p_pill.space_after = Pt(12)

    p_title = h_tf.add_paragraph()
    p_title.text = "DevAssist AI"
    p_title.font.name = FONT_TITLE
    p_title.font.size = Pt(46)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE
    p_title.alignment = PP_ALIGN.CENTER
    p_title.space_after = Pt(10)

    p_sub = h_tf.add_paragraph()
    p_sub.text = "Fast Bug Detection, Full Project Understanding, and 1-Click Fixes"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(16)
    p_sub.font.bold = False
    p_sub.font.color.rgb = ACCENT_PRIMARY
    p_sub.alignment = PP_ALIGN.CENTER

    # 4 Meta Information Cards
    meta_info = [
        ("TEAM MEMBERS", "Atharv Jain\nYashvardhan", ACCENT_CYAN),
        ("UNIVERSITY", "BML Munjal\nUniversity", ACCENT_PRIMARY),
        ("COURSE / SUBJECT", "AI DevOps\n(Automated Review)", ACCENT_PURPLE),
        ("FACULTY MENTOR", "Mr. Naveen Singh\nAcademic Advisor", ACCENT_EMERALD),
    ]

    card_w = Inches(2.55)
    gap = Inches(0.24)
    start_x = Inches(1.2)
    top_y = Inches(5.3)

    for i, (label, val, accent) in enumerate(meta_info):
        cx = start_x + i * (card_w + gap)
        mtf = create_card_with_text(s1, cx, top_y, card_w, Inches(1.5), bg_color=CARD_BG, border_color=CARD_BORDER, accent_bar=accent, pad_x=0.18, pad_y=0.14)
        
        mp0 = mtf.paragraphs[0]
        mp0.text = label
        mp0.font.name = FONT_BODY
        mp0.font.size = Pt(10)
        mp0.font.bold = True
        mp0.font.color.rgb = accent
        mp0.space_after = Pt(4)

        mp1 = mtf.add_paragraph()
        mp1.text = val
        mp1.font.name = FONT_TITLE
        mp1.font.size = Pt(12)
        mp1.font.bold = True
        mp1.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 2: Problem → Aim
    # =========================================================================
    s2 = create_blank_slide(prs)
    add_header(s2, "Problem → Goal: Why Do Teams Need DevAssist AI?", "BACKGROUND & PURPOSE", 2)

    # Left: The Problem Card
    ptf = create_card_with_text(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.4), accent_bar=ACCENT_ROSE, pad_x=0.25, pad_y=0.2)
    p_h = ptf.paragraphs[0]
    p_h.text = "The Real Problem"
    p_h.font.name = FONT_TITLE
    p_h.font.size = Pt(16)
    p_h.font.bold = True
    p_h.font.color.rgb = TEXT_WHITE
    p_h.space_after = Pt(2)

    p_sub = ptf.add_paragraph()
    p_sub.text = "Manual Code Reviews Are Slow, Tiring, and Miss Hidden Bugs"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(10.5)
    p_sub.font.color.rgb = ACCENT_ROSE
    p_sub.space_after = Pt(14)

    problems = [
        ("Long Review Delays", "Senior developers spend 20%+ of their day reviewing code line-by-line, causing major delays in releasing features."),
        ("Too Many Programming Languages", "Modern projects combine Python, JavaScript, Go, and SQL. Nobody is an expert in every language."),
        ("Hidden Ripple Effects", "Changing a single shared function in one file can quietly break 10+ other files without anyone realizing."),
        ("Wasted Time on Tiny Details", "Engineers spend hours debating formatting, spacing, and commas instead of finding critical security holes."),
        ("Noisy, Fake AI Warnings", "Generic AI tools spam pull requests with false alarms, wrong line numbers, and unhelpful advice."),
    ]
    for title, desc in problems:
        pt = ptf.add_paragraph()
        pt.text = f"❌  {title}"
        pt.font.name = FONT_TITLE
        pt.font.size = Pt(11)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_WHITE
        pt.space_after = Pt(1)

        pd = ptf.add_paragraph()
        pd.text = desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(7)

    # Right: The Aim Card
    atf = create_card_with_text(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(2.2), accent_bar=ACCENT_EMERALD, pad_x=0.25, pad_y=0.18)
    a_h = atf.paragraphs[0]
    a_h.text = "Our Goal & Mission"
    a_h.font.name = FONT_TITLE
    a_h.font.size = Pt(16)
    a_h.font.bold = True
    a_h.font.color.rgb = TEXT_WHITE
    a_h.space_after = Pt(2)

    a_sub = atf.add_paragraph()
    a_sub.text = "Fast, Accurate, and Trustworthy AI Code Reviews"
    a_sub.font.name = FONT_BODY
    a_sub.font.size = Pt(10.5)
    a_sub.font.color.rgb = ACCENT_EMERALD
    a_sub.space_after = Pt(8)

    a_body = atf.add_paragraph()
    a_body.text = "DevAssist AI automates the code review process. When a developer creates a pull request, the AI reads the changes, understands the entire project, spots dangerous bugs, suggests tested 1-click fixes, and writes clear summaries—saving engineering teams hours every day."
    a_body.font.name = FONT_BODY
    a_body.font.size = Pt(9.5)
    a_body.font.color.rgb = TEXT_MUTED
    a_body.line_spacing = 1.2

    # Right Bottom: 4 Distinct Pillars (2x2 Grid)
    pillars = [
        ("⚡ FAST REVIEWS", "Delivers complete, thorough reviews in under 60 seconds so developers never wait.", ACCENT_CYAN),
        ("🧠 3-STEP AI CHECKS", "Summarizes, finds bugs, and verifies accuracy with multiple checks to prevent mistakes.", ACCENT_PRIMARY),
        ("🔍 FULL PROJECT CONTEXT", "Traces how files connect together so changes in one place don't break other features.", ACCENT_PURPLE),
        ("👤 HUMAN STAYS IN CHARGE", "The AI gives smart suggestions, but human developers always have the final say to merge.", ACCENT_EMERALD),
    ]

    for idx, (p_title, p_desc, p_color) in enumerate(pillars):
        px = Inches(6.8) if idx % 2 == 0 else Inches(9.75)
        py = Inches(4.0) if idx < 2 else Inches(5.45)
        pw = Inches(2.75)
        ph = Inches(1.35)

        pctf = create_card_with_text(s2, px, py, pw, ph, bg_color=CARD_BG_HOVER, pad_x=0.18, pad_y=0.14)
        pp0 = pctf.paragraphs[0]
        pp0.text = p_title
        pp0.font.name = FONT_BODY
        pp0.font.size = Pt(10.5)
        pp0.font.bold = True
        pp0.font.color.rgb = p_color
        pp0.space_after = Pt(3)

        pp1 = pctf.add_paragraph()
        pp1.text = p_desc
        pp1.font.name = FONT_BODY
        pp1.font.size = Pt(9)
        pp1.font.color.rgb = TEXT_MUTED
        pp1.line_spacing = 1.15

    # =========================================================================
    # SLIDE 3: Multi-Agent & Multi-LLM Review System
    # =========================================================================
    s3 = create_blank_slide(prs)
    add_header(s3, "Our Approach: A Team of Specialized AI Assistants", "HOW THE SYSTEM WORKS", 3)

    agents = [
        ("1. Change Summarizer", "Fast AI (Gemini)", "Quickly reads the pull request changes and creates a neat, easy summary of what was modified.", ACCENT_CYAN),
        ("2. Deep Bug Hunter", "Deep AI (Claude 3.5)", "Carefully inspects the code for logic errors, security flaws, crashes, and drafts working fixes.", ACCENT_PRIMARY),
        ("3. Accuracy Checker", "Fast AI (Gemini)", "Double-checks every finding. Automatically drops weak or uncertain warnings to prevent spam.", ACCENT_EMERALD),
        ("4. Dependency Map", "CodeGraph Parser", "Traces which functions call each other so you know what other parts of the project might break.", ACCENT_PURPLE),
        ("5. Smart Code Search", "Hybrid Search (RAG)", "Instantly finds related functions and data models from anywhere in the codebase for full context.", ACCENT_CYAN),
        ("6. Standard Linters", "Ruff, Pylint, ESLint", "Runs standard industry tools to catch syntax errors, unused imports, and styling mistakes.", ACCENT_AMBER),
        ("7. Safe Test Sandbox", "Isolated Docker Box", "Runs dynamic validation tests inside a secure locked box without any risk to the main computer.", ACCENT_ROSE),
        ("8. Chat Assistant", "Multi-AI Router", "Developers can reply to any comment on GitHub, and the AI answers follow-up questions directly.", ACCENT_PRIMARY),
        ("9. Auto-Documentation", "Doc Writer", "Automatically writes clean documentation and code comments once a pull request is approved.", ACCENT_EMERALD),
    ]

    gw = Inches(3.7)
    gh = Inches(1.6)
    gx_start = Inches(0.8)
    gy_start = Inches(1.6)
    x_gap = Inches(0.3)
    y_gap = Inches(0.25)

    for i, (a_name, a_model, a_desc, a_acc) in enumerate(agents):
        row = i // 3
        col = i % 3
        x = gx_start + col * (gw + x_gap)
        y = gy_start + row * (gh + y_gap)

        actf = create_card_with_text(s3, x, y, gw, gh, accent_bar=a_acc, pad_x=0.2, pad_y=0.14)

        ap0 = actf.paragraphs[0]
        ap0.text = a_name
        ap0.font.name = FONT_TITLE
        ap0.font.size = Pt(12)
        ap0.font.bold = True
        ap0.font.color.rgb = TEXT_WHITE
        ap0.space_after = Pt(2)

        ap1 = actf.add_paragraph()
        ap1.text = f"Powered by: {a_model}"
        ap1.font.name = FONT_MONO
        ap1.font.size = Pt(9.5)
        ap1.font.bold = True
        ap1.font.color.rgb = a_acc
        ap1.space_after = Pt(4)

        ap2 = actf.add_paragraph()
        ap2.text = a_desc
        ap2.font.name = FONT_BODY
        ap2.font.size = Pt(9)
        ap2.font.color.rgb = TEXT_MUTED
        ap2.line_spacing = 1.15

    # =========================================================================
    # SLIDE 4: How DevAssist AI Works – End-to-End
    # =========================================================================
    s4 = create_blank_slide(prs)
    add_header(s4, "Step-by-Step: How a Code Review Happens", "STEP-BY-STEP PROCESS", 4)

    steps = [
        ("1. PR is Opened", "GitHub or GitLab", "A developer submits new code changes for review.", ACCENT_CYAN),
        ("2. Instant Pickup", "Fast Task Queue", "DevAssist AI immediately starts a background task in seconds.", ACCENT_PRIMARY),
        ("3. Map Connections", "Dependency Scanner", "Checks which other files and functions touch this code.", ACCENT_PURPLE),
        ("4. Gather Context", "Smart Code Search", "Pulls related helper functions and types across the repo.", ACCENT_EMERALD),
        ("5. Summarize Scope", "Fast AI Summary", "Creates a concise overview of what this pull request changes.", ACCENT_CYAN),
        ("6. Find Real Bugs", "Deep Logic Inspection", "Spots security vulnerabilities, logic flaws, and drafts fixes.", ACCENT_PRIMARY),
        ("7. Verify Findings", "Accuracy Filter", "Drops low-confidence warnings and verifies exact line numbers.", ACCENT_AMBER),
        ("8. Post on GitHub", "Direct PR Comments", "Posts inline suggestions, quality scores, and full summary.", ACCENT_EMERALD),
    ]

    sw = Inches(2.75)
    sh = Inches(1.7)
    sx_start = Inches(0.8)
    sy_start = Inches(1.6)
    sgap = Inches(0.24)

    for i, (stitle, ssub, sbody, sacc) in enumerate(steps):
        row = i // 4
        col = i % 4
        x = sx_start + col * (sw + sgap)
        y = sy_start + row * (sh + Inches(0.25))

        ctf = create_card_with_text(s4, x, y, sw, sh, accent_bar=sacc, pad_x=0.18, pad_y=0.14)

        cp0 = ctf.paragraphs[0]
        cp0.text = stitle
        cp0.font.name = FONT_TITLE
        cp0.font.size = Pt(12)
        cp0.font.bold = True
        cp0.font.color.rgb = TEXT_WHITE
        cp0.space_after = Pt(2)

        cp1 = ctf.add_paragraph()
        cp1.text = ssub
        cp1.font.name = FONT_BODY
        cp1.font.size = Pt(9)
        cp1.font.bold = True
        cp1.font.color.rgb = sacc
        cp1.space_after = Pt(4)

        cp2 = ctf.add_paragraph()
        cp2.text = sbody
        cp2.font.name = FONT_BODY
        cp2.font.size = Pt(8.5)
        cp2.font.color.rgb = TEXT_MUTED
        cp2.line_spacing = 1.15

    # Bottom 3 Architecture Pillars
    bottom_cards = [
        ("Understands Real Code Structure", "Reads complete functions and classes as whole units, not chopped-up random text lines.", ACCENT_CYAN),
        ("Smooth Background Speed", "Runs all reviews in the background so developers never experience delays or frozen screens.", ACCENT_PRIMARY),
        ("Live Progress Updates", "Watch the review advance step-by-step on the web dashboard in real-time as it finishes.", ACCENT_EMERALD),
    ]

    bw = Inches(3.7)
    bh = Inches(1.4)
    by = Inches(5.55)

    for i, (btitle, bdesc, bacc) in enumerate(bottom_cards):
        bx = sx_start + i * (bw + Inches(0.31))
        bctf = create_card_with_text(s4, bx, by, bw, bh, bg_color=CARD_BG_HOVER, pad_x=0.18, pad_y=0.14)

        bp0 = bctf.paragraphs[0]
        bp0.text = btitle
        bp0.font.name = FONT_TITLE
        bp0.font.size = Pt(11)
        bp0.font.bold = True
        bp0.font.color.rgb = bacc
        bp0.space_after = Pt(3)

        bp1 = bctf.add_paragraph()
        bp1.text = bdesc
        bp1.font.name = FONT_BODY
        bp1.font.size = Pt(9)
        bp1.font.color.rgb = TEXT_MUTED
        bp1.line_spacing = 1.15

    # =========================================================================
    # SLIDE 5: From Findings to Decisions
    # =========================================================================
    s5 = create_blank_slide(prs)
    add_header(s5, "How We Ensure Quality: Filtering Bugs & Scoring PRs", "VERIFICATION & SCORING", 5)

    col_w = Inches(3.7)
    col_h = Inches(5.4)
    cy = Inches(1.6)

    # Col 1: Ripple Effect (Blast Radius)
    c1_tf = create_card_with_text(s5, Inches(0.8), cy, col_w, col_h, accent_bar=ACCENT_ROSE, pad_x=0.22, pad_y=0.18)
    p0 = c1_tf.paragraphs[0]
    p0.text = "1. Ripple Effect Analysis"
    p0.font.name = FONT_TITLE
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_WHITE
    p0.space_after = Pt(2)

    p_sub = c1_tf.add_paragraph()
    p_sub.text = "Seeing What Else Could Break"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(10)
    p_sub.font.color.rgb = ACCENT_ROSE
    p_sub.space_after = Pt(12)

    c1_pts = [
        ("Changed Code", "Example: A login or authentication function is modified."),
        ("Tracing Callers", "The system scans all files to see what calls that function."),
        ("Affected Endpoints", "Finds all 14 API routes and screens that depend on it."),
        ("Risk Warning", "Flags a 'High Ripple Effect' so engineers test carefully."),
        ("Real Benefit", "Stops unexpected breakages before code reaches users."),
    ]
    for h, b in c1_pts:
        p = c1_tf.add_paragraph()
        p.text = f"• {h}"
        p.font.name = FONT_TITLE
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(1)

        pd = c1_tf.add_paragraph()
        pd.text = b
        pd.font.name = FONT_BODY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(6)

    # Col 2: The Accuracy Filter
    c2_tf = create_card_with_text(s5, Inches(4.8), cy, col_w, col_h, accent_bar=ACCENT_AMBER, pad_x=0.22, pad_y=0.18)
    p0 = c2_tf.paragraphs[0]
    p0.text = "2. The Accuracy Filter"
    p0.font.name = FONT_TITLE
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_WHITE
    p0.space_after = Pt(2)

    p_sub = c2_tf.add_paragraph()
    p_sub.text = "Stopping False Alarms & Hallucinations"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(10)
    p_sub.font.color.rgb = ACCENT_AMBER
    p_sub.space_after = Pt(12)

    c2_pts = [
        ("Correct Line Numbers", "Makes sure the reported issue is actually on the changed lines."),
        ("Working Code Fixes", "Verifies that suggested code fixes have valid, working syntax."),
        ("Confidence Rating", "Rates how certain the AI is about the bug from 0% to 100%."),
        ("Drops Weak Alerts", "Throws away any warning that scores below 60% confidence."),
        ("No Repeated Spam", "Never re-posts the same comment twice across updates."),
    ]
    for h, b in c2_pts:
        p = c2_tf.add_paragraph()
        p.text = f"✔ {h}"
        p.font.name = FONT_TITLE
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = ACCENT_EMERALD
        p.space_after = Pt(1)

        pd = c2_tf.add_paragraph()
        pd.text = b
        pd.font.name = FONT_BODY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(6)

    # Col 3: PR Quality Score
    c3_tf = create_card_with_text(s5, Inches(8.8), cy, col_w, col_h, accent_bar=ACCENT_EMERALD, pad_x=0.22, pad_y=0.18)
    p0 = c3_tf.paragraphs[0]
    p0.text = "3. Simple PR Quality Score"
    p0.font.name = FONT_TITLE
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_WHITE
    p0.space_after = Pt(2)

    p_sub = c3_tf.add_paragraph()
    p_sub.text = "A Clear 0 to 10 Health Rating"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(10)
    p_sub.font.color.rgb = ACCENT_EMERALD
    p_sub.space_after = Pt(12)

    c3_pts = [
        ("🔴 Critical Bugs (-2.0 pts)", "Security risks, data corruption, app crashes."),
        ("🟡 Warnings (-0.5 pts)", "Missing timeouts, slow queries, unhandled errors."),
        ("🔵 Helpful Notes (-0.1 pts)", "Missing comments, naming cleanup, minor polish."),
        ("Clear Score (0 to 10)", "Example: 🟢 8.5/10 lets the team know code health at a glance."),
        ("Developer in Control", "Engineers can accept fixes, dismiss notes, or merge."),
    ]
    for h, b in c3_pts:
        p = c3_tf.add_paragraph()
        p.text = h
        p.font.name = FONT_TITLE
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(1)

        pd = c3_tf.add_paragraph()
        pd.text = b
        pd.font.name = FONT_BODY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(6)

    # =========================================================================
    # SLIDE 6: System Architecture & Technology Stack
    # =========================================================================
    s6 = create_blank_slide(prs)
    add_header(s6, "Under the Hood: Technologies Powering DevAssist AI", "MODERN & RELIABLE TECH STACK", 6)

    tech_pillars = [
        ("Modern Web Dashboard", "Clean & Responsive UI", [
            ("Framework", "Next.js 16 & React 19 (Fast page loading)"),
            ("Code Diff Viewer", "Interactive side-by-side view with highlighting"),
            ("Visual Diagrams", "Interactive node maps showing file connections"),
            ("Flowcharts", "Live-rendered Mermaid system flowcharts"),
            ("User Login", "Secure user authentication and workspace control"),
        ], ACCENT_CYAN),

        ("Fast Async Backend", "Built for High Performance", [
            ("API Server", "FastAPI (Lightweight, lightning-fast Python API)"),
            ("Task Queue", "Celery & Redis (Runs heavy reviews in background)"),
            ("Real-Time Events", "Live streaming from server to browser without lag"),
            ("Database", "PostgreSQL & SQLite (Saves PRs, reviews, & history)"),
            ("Auto-Updates", "Automated database tools prevent data loss"),
        ], ACCENT_PRIMARY),

        ("Multi-AI Brain", "Reliable Multi-Model System", [
            ("Deep Reasoning", "Claude 3.5 Sonnet & GPT-4o for complex bug finding"),
            ("Fast Summaries", "Google Gemini 2.0 Flash for instant reading"),
            ("Backup Chain", "Auto-switches to backup AI if primary model fails"),
            ("Smart Search", "High-speed vector database (FAISS) + keyword search"),
            ("Best-Match Fusion", "Combines keywords and meaning for accurate context"),
        ], ACCENT_PURPLE),

        ("Safe Infrastructure", "Containerized & Protected", [
            ("Containers", "Docker & Docker Compose (Runs anywhere in 1 click)"),
            ("Safe Sandbox", "Runs test scripts inside an isolated locked container"),
            ("Zero-Risk Security", "No internet access, cannot modify or damage host"),
            ("Git Integration", "Direct webhooks for GitHub and GitLab workflows"),
            ("Industry Linters", "Runs Ruff, ESLint, Bandit, and Mypy in parallel"),
        ], ACCENT_EMERALD),
    ]

    pw = Inches(2.75)
    ph = Inches(5.4)
    px_start = Inches(0.8)
    py_start = Inches(1.6)
    pgap = Inches(0.24)

    for i, (t_title, t_sub, items, t_acc) in enumerate(tech_pillars):
        x = px_start + i * (pw + pgap)
        itf = create_card_with_text(s6, x, py_start, pw, ph, accent_bar=t_acc, pad_x=0.18, pad_y=0.16)

        p0 = itf.paragraphs[0]
        p0.text = t_title
        p0.font.name = FONT_TITLE
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_WHITE
        p0.space_after = Pt(2)

        p1 = itf.add_paragraph()
        p1.text = t_sub
        p1.font.name = FONT_BODY
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = t_acc
        p1.space_after = Pt(10)

        for k, v in items:
            pk = itf.add_paragraph()
            pk.text = k
            pk.font.name = FONT_TITLE
            pk.font.size = Pt(9.5)
            pk.font.bold = True
            pk.font.color.rgb = t_acc
            pk.space_after = Pt(1)

            pv = itf.add_paragraph()
            pv.text = v
            pv.font.name = FONT_BODY
            pv.font.size = Pt(8.5)
            pv.font.color.rgb = TEXT_WHITE
            pv.space_after = Pt(5)

    # =========================================================================
    # SLIDE 7: What the User Sees
    # =========================================================================
    s7 = create_blank_slide(prs)
    add_header(s7, "Where Developers Use It: Web Dashboard & GitHub", "USER EXPERIENCE & INTERACTION", 7)

    # Left: Web Dashboard
    dtf = create_card_with_text(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.4), accent_bar=ACCENT_CYAN, pad_x=0.25, pad_y=0.2)
    dh = dtf.paragraphs[0]
    dh.text = "Web Dashboard Experience"
    dh.font.name = FONT_TITLE
    dh.font.size = Pt(16)
    dh.font.bold = True
    dh.font.color.rgb = TEXT_WHITE
    dh.space_after = Pt(2)

    dsub = dtf.add_paragraph()
    dsub.text = "Central Command Center for Your Projects"
    dsub.font.name = FONT_BODY
    dsub.font.size = Pt(10.5)
    dsub.font.color.rgb = ACCENT_CYAN
    dsub.space_after = Pt(14)

    dash_features = [
        ("Project Health Dashboard", "See review speed, bug trends, and overall code quality over time."),
        ("Visual File Connection Map", "Interactive diagram showing which files were modified and what else they affect."),
        ("Chapter-by-Chapter Review", "Splits big 50-file pull requests into easy logical chapters so reviews aren't overwhelming."),
        ("System Architecture Diagrams", "Automatically generates visual flowcharts showing how data moves through the new code."),
        ("Ask AI About Your Code", "Search the entire codebase in plain English to find functions, models, and explanations."),
    ]
    for f_title, f_desc in dash_features:
        p = dtf.add_paragraph()
        p.text = f"📊  {f_title}"
        p.font.name = FONT_TITLE
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        p.space_after = Pt(1)

        pd = dtf.add_paragraph()
        pd.text = f_desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(8)

    # Right: Native GitHub PR Experience
    ghtf = create_card_with_text(s7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.4), accent_bar=ACCENT_PRIMARY, pad_x=0.25, pad_y=0.2)
    gh_h = ghtf.paragraphs[0]
    gh_h.text = "GitHub & GitLab Experience"
    gh_h.font.name = FONT_TITLE
    gh_h.font.size = Pt(16)
    gh_h.font.bold = True
    gh_h.font.color.rgb = TEXT_WHITE
    gh_h.space_after = Pt(2)

    gh_sub = ghtf.add_paragraph()
    gh_sub.text = "Built Right Into Your Normal Daily Workflow"
    gh_sub.font.name = FONT_BODY
    gh_sub.font.size = Pt(10.5)
    gh_sub.font.color.rgb = ACCENT_PRIMARY
    gh_sub.space_after = Pt(14)

    gh_features = [
        ("Line-by-Line Review Comments", "Helpful, clear explanations right on the exact lines where issues are found."),
        ("1-Click Code Fixes (```suggestion)", "Click 'Apply Suggestion' on GitHub to instantly accept and commit the fix without typing."),
        ("Clear Summary Walkthrough", "A neat table at the top of the PR summarizing all changes, files, and impact."),
        ("Visual Quality Scorecard", "A simple badge (e.g., 🟢 9.0/10) showing PR readiness and bug counts."),
        ("Chat Directly on Comments", "Reply to any bot comment to ask 'Why is this better?' or request another solution."),
    ]
    for g_title, g_desc in gh_features:
        p = ghtf.add_paragraph()
        p.text = f"🤖  {g_title}"
        p.font.name = FONT_TITLE
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_PRIMARY
        p.space_after = Pt(1)

        pd = ghtf.add_paragraph()
        pd.text = g_desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(9)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(8)

    # =========================================================================
    # SLIDE 8: Where AI Reduces Effort
    # =========================================================================
    s8 = create_blank_slide(prs)
    add_header(s8, "Before vs After: How DevAssist AI Saves Time", "REAL PRODUCTIVITY COMPARISON", 8)

    # Left & Right Headers
    h_man_tf = create_card_with_text(s8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(0.52), bg_color=RGBColor(36, 18, 26), border_color=ACCENT_ROSE, pad_x=0.15, pad_y=0.12)
    mp = h_man_tf.paragraphs[0]
    mp.text = "❌  Traditional Manual Review"
    mp.font.name = FONT_TITLE
    mp.font.size = Pt(11.5)
    mp.font.bold = True
    mp.font.color.rgb = ACCENT_ROSE

    h_ai_tf = create_card_with_text(s8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(0.52), bg_color=RGBColor(16, 36, 30), border_color=ACCENT_EMERALD, pad_x=0.15, pad_y=0.12)
    ap = h_ai_tf.paragraphs[0]
    ap.text = "⚡  With DevAssist AI"
    ap.font.name = FONT_TITLE
    ap.font.size = Pt(11.5)
    ap.font.bold = True
    ap.font.color.rgb = ACCENT_EMERALD

    comparisons = [
        ("Reading raw code line-by-line", "Developers read isolated code diffs without knowing helper functions in other files.",
         "Instant full project context", "The AI automatically pulls in related files and helper functions for full context."),

        ("Missing hidden ripple effects", "Accidental breakages in other files are discovered only after staging or production crashes.",
         "Automatic connection mapping", "Instantly maps all callers and warns if changing a function breaks other parts of the project."),

        ("Arguing over formatting and commas", "Reviewers waste time pointing out indentation and missing imports instead of core logic.",
         "Automated style cleanup", "Automated linters catch style issues immediately, keeping human focus on real bugs."),

        ("Spammed by noisy AI warnings", "Basic AI bots post dozens of messy comments, many of which are fake or unhelpful.",
         "Strict accuracy filtering", "A verification filter drops weak warnings, keeping only high-confidence, helpful feedback."),

        ("Outdated or forgotten documentation", "Engineers rush to ship features, leaving READMEs and code comments outdated.",
         "Auto-written documentation", "The AI writes clean documentation and comments automatically as soon as code is merged."),
    ]

    row_y = Inches(2.25)
    row_h = Inches(0.85)
    rgap = Inches(0.12)

    for i, (m_title, m_desc, a_title, a_desc) in enumerate(comparisons):
        y = row_y + i * (row_h + rgap)

        mctf = create_card_with_text(s8, Inches(0.8), y, Inches(5.6), row_h, pad_x=0.16, pad_y=0.1)
        p0 = mctf.paragraphs[0]
        p0.text = m_title
        p0.font.name = FONT_TITLE
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_WHITE
        p0.space_after = Pt(1)

        p1 = mctf.add_paragraph()
        p1.text = m_desc
        p1.font.name = FONT_BODY
        p1.font.size = Pt(8.5)
        p1.font.color.rgb = TEXT_MUTED

        actf = create_card_with_text(s8, Inches(6.8), y, Inches(5.7), row_h, pad_x=0.16, pad_y=0.1)
        ap0 = actf.paragraphs[0]
        ap0.text = a_title
        ap0.font.name = FONT_TITLE
        ap0.font.size = Pt(10)
        ap0.font.bold = True
        ap0.font.color.rgb = ACCENT_EMERALD
        ap0.space_after = Pt(1)

        ap1 = actf.add_paragraph()
        ap1.text = a_desc
        ap1.font.name = FONT_BODY
        ap1.font.size = Pt(8.5)
        ap1.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 9: What Improves with DevAssist AI? (Metrics)
    # =========================================================================
    s9 = create_blank_slide(prs)
    add_header(s9, "Key Results: Why Teams Love DevAssist AI", "MEASURABLE IMPACT & BENEFITS", 9)

    metrics = [
        ("70%+", "FASTER REVIEWS", "Background processing completes thorough reviews in under 60 seconds instead of hours of waiting.", ACCENT_CYAN),
        ("< 5%", "FALSE ALARM RATE", "Multi-step verification removes noisy hallucinations, giving developers advice they can actually trust.", ACCENT_EMERALD),
        ("3 LEVELS", "DEPENDENCY DEPTH", "Traces function connections up to 3 levels deep across the whole project to catch breaking changes before merge.", ACCENT_PRIMARY),
        ("99.9%", "RELIABLE UPTIME", "If one AI provider is down or slow, the system automatically switches to a backup model without failing.", ACCENT_PURPLE),
        ("1-CLICK", "INSTANT CODE FIXES", "Provides ready-to-use code suggestions that developers can accept directly inside GitHub with one click.", ACCENT_AMBER),
    ]

    mw = Inches(2.18)
    mh = Inches(5.2)
    mx_start = Inches(0.8)
    my = Inches(1.7)
    mgap = Inches(0.2)

    for i, (stat, stat_lbl, stat_desc, stat_acc) in enumerate(metrics):
        x = mx_start + i * (mw + mgap)
        ctf = create_card_with_text(s9, x, my, mw, mh, accent_bar=stat_acc, pad_x=0.16, pad_y=0.25)

        p0 = ctf.paragraphs[0]
        p0.text = stat
        p0.font.name = FONT_TITLE
        p0.font.size = Pt(32)
        p0.font.bold = True
        p0.font.color.rgb = stat_acc
        p0.alignment = PP_ALIGN.CENTER
        p0.space_after = Pt(4)

        p1 = ctf.add_paragraph()
        p1.text = stat_lbl
        p1.font.name = FONT_BODY
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.alignment = PP_ALIGN.CENTER
        p1.space_after = Pt(16)

        p2 = ctf.add_paragraph()
        p2.text = stat_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_MUTED
        p2.line_spacing = 1.25

    # =========================================================================
    # SLIDE 10: DevOps & Deployment
    # =========================================================================
    s10 = create_blank_slide(prs)
    add_header(s10, "How It Runs: Safe, Reliable, and Easy to Deploy", "DEPLOYMENT & SECURITY", 10)

    devops_steps = [
        ("1. Push Code", "Push to GitHub or GitLab.", ACCENT_CYAN),
        ("2. Build Image", "Package in Docker containers.", ACCENT_PRIMARY),
        ("3. Start Services", "Run Redis, DB, API & worker.", ACCENT_PURPLE),
        ("4. Webhook Ingress", "Securely receive PR events.", ACCENT_AMBER),
        ("5. Safe Sandbox", "Run tests in locked container.", ACCENT_ROSE),
        ("6. Live Streaming", "Stream updates to dashboard.", ACCENT_EMERALD),
    ]

    dw = Inches(1.8)
    dh = Inches(1.05)
    dx_start = Inches(0.8)
    dy = Inches(1.6)
    dgap = Inches(0.18)

    for i, (d_title, d_sub, d_acc) in enumerate(devops_steps):
        x = dx_start + i * (dw + dgap)
        ctf = create_card_with_text(s10, x, dy, dw, dh, accent_bar=d_acc, pad_x=0.12, pad_y=0.1)

        cp0 = ctf.paragraphs[0]
        cp0.text = d_title
        cp0.font.name = FONT_TITLE
        cp0.font.size = Pt(10)
        cp0.font.bold = True
        cp0.font.color.rgb = TEXT_WHITE
        cp0.space_after = Pt(2)

        cp1 = ctf.add_paragraph()
        cp1.text = d_sub
        cp1.font.name = FONT_BODY
        cp1.font.size = Pt(8.5)
        cp1.font.color.rgb = TEXT_MUTED

    devops_cards = [
        ("Complete Docker Setup", "One Simple Command to Start Everything", [
            ("Redis 7", "Handles fast task queues and real-time live event streaming."),
            ("PostgreSQL 15", "Securely stores repositories, review logs, and bug history."),
            ("Celery Worker", "Runs heavy AI reviews smoothly in the background without lag."),
            ("FastAPI Server", "Lightning-fast API server with secure GitHub webhook verification."),
        ], ACCENT_CYAN),

        ("Safe & Secure Testing Sandbox", "Zero Risk to Your Host Machines", [
            ("Isolated Box", "Test scripts run in a locked container with zero internet access."),
            ("Locked Permissions", "Cannot access host files or sensitive operating system data."),
            ("Read-Only Code", "Cloned code cannot be corrupted or accidentally modified."),
            ("Auto-Stop Timeout", "Automatically stops any script taking longer than 30 seconds."),
        ], ACCENT_ROSE),

        ("Monitoring & Reliability", "Always Know System Health", [
            ("Live Health Checks", "Constant monitoring of database and AI service connectivity."),
            ("Live UI Progress", "Watch reviews advance step-by-step on the web dashboard."),
            ("Clean DB Updates", "Automated database tools prevent downtime during updates."),
            ("Complete Audit Trail", "Tracks every review, comment, and developer approval."),
        ], ACCENT_EMERALD),
    ]

    cw = Inches(3.7)
    ch = Inches(4.2)
    cy = Inches(2.85)

    for i, (title, sub, bullet_items, acc) in enumerate(devops_cards):
        cx = Inches(0.8) + i * (cw + Inches(0.31))
        ctf = create_card_with_text(s10, cx, cy, cw, ch, accent_bar=acc, pad_x=0.2, pad_y=0.18)

        p0 = ctf.paragraphs[0]
        p0.text = title
        p0.font.name = FONT_TITLE
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_WHITE
        p0.space_after = Pt(2)

        p1 = ctf.add_paragraph()
        p1.text = sub
        p1.font.name = FONT_BODY
        p1.font.size = Pt(9.5)
        p1.font.color.rgb = acc
        p1.space_after = Pt(10)

        for k, v in bullet_items:
            pk = ctf.add_paragraph()
            pk.text = f"• {k}"
            pk.font.name = FONT_TITLE
            pk.font.size = Pt(9.5)
            pk.font.bold = True
            pk.font.color.rgb = acc
            pk.space_after = Pt(1)

            pv = ctf.add_paragraph()
            pv.text = v
            pv.font.name = FONT_BODY
            pv.font.size = Pt(8.5)
            pv.font.color.rgb = TEXT_MUTED
            pv.space_after = Pt(4)

    # =========================================================================
    # SLIDE 11: Future Scope
    # =========================================================================
    s11 = create_blank_slide(prs)
    add_header(s11, "What's Next: Future Roadmap & Upcoming Features", "FUTURE PLANS & EXPANSION", 11)

    future_items = [
        ("In-Editor Plugins", "VS Code & JetBrains Tools", "Get instant AI review feedback right inside your code editor before you even push your code.", ACCENT_CYAN),
        ("Self-Healing Bug Fixes", "Automated Fix Branches", "For simple bugs, the AI creates a dedicated fix branch with verified tests ready for 1-click merge.", ACCENT_PRIMARY),
        ("Cross-Project Knowledge", "Multi-Repo Support", "Connect multiple related projects together to verify shared APIs and microservices.", ACCENT_PURPLE),
        ("Custom Team Rules", "Company Coding Standards", "Teach the AI your team's specific coding guidelines, security policies, and naming conventions.", ACCENT_EMERALD),
        ("Team Permissions & Compliance", "Enterprise Access Control", "Add role-based permissions, team workspaces, and exportable security audit reports.", ACCENT_AMBER),
        ("Automatic Test Generation", "AI-Written Unit Tests", "Automatically write unit and integration tests covering newly added code to prevent future regressions.", ACCENT_ROSE),
    ]

    fw = Inches(3.7)
    fh = Inches(2.25)
    fx_start = Inches(0.8)
    fy_start = Inches(1.6)
    fgap_x = Inches(0.31)
    fgap_y = Inches(0.25)

    for i, (f_title, f_sub, f_desc, f_acc) in enumerate(future_items):
        row = i // 3
        col = i % 3
        x = fx_start + col * (fw + fgap_x)
        y = fy_start + row * (fh + fgap_y)

        ctf = create_card_with_text(s11, x, y, fw, fh, accent_bar=f_acc, pad_x=0.2, pad_y=0.14)

        p0 = ctf.paragraphs[0]
        p0.text = f_title
        p0.font.name = FONT_TITLE
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_WHITE
        p0.space_after = Pt(2)

        p1 = ctf.add_paragraph()
        p1.text = f_sub
        p1.font.name = FONT_BODY
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = f_acc
        p1.space_after = Pt(6)

        p2 = ctf.add_paragraph()
        p2.text = f_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9)
        p2.font.color.rgb = TEXT_MUTED
        p2.line_spacing = 1.15

    # Bottom Goal Banner
    mbtf = create_card_with_text(s11, Inches(0.8), Inches(6.45), Inches(11.733), Inches(0.65), bg_color=CARD_BG_HOVER, border_color=ACCENT_PRIMARY, pad_x=0.2, pad_y=0.12)
    mp = mbtf.paragraphs[0]
    mp.text = "🎯  Our Goal: Make DevAssist AI the smartest, safest, and most helpful code review partner for every software team."
    mp.font.name = FONT_BODY
    mp.font.size = Pt(10)
    mp.font.bold = True
    mp.font.color.rgb = TEXT_WHITE
    mp.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 12: Thank You & Q&A
    # =========================================================================
    s12 = create_blank_slide(prs)

    # Hero Box
    hero12 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.0), Inches(10.333), Inches(5.4))
    hero12.fill.solid()
    hero12.fill.fore_color.rgb = CARD_BG
    hero12.line.color.rgb = ACCENT_PRIMARY
    hero12.line.width = Pt(1.5)

    ostrip = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.0), Inches(10.333), Inches(0.08))
    ostrip.fill.solid()
    ostrip.fill.fore_color.rgb = ACCENT_CYAN
    ostrip.line.fill.background()

    # Thank you texts in single text frame
    ot_box = s12.shapes.add_textbox(Inches(2.0), Inches(1.4), Inches(9.333), Inches(2.4))
    otf = ot_box.text_frame
    otf.word_wrap = True
    otf.margin_left = 0
    otf.margin_top = 0
    otf.margin_right = 0
    otf.margin_bottom = 0

    op0 = otf.paragraphs[0]
    op0.text = "Thank You!"
    op0.font.name = FONT_TITLE
    op0.font.size = Pt(44)
    op0.font.bold = True
    op0.font.color.rgb = TEXT_WHITE
    op0.alignment = PP_ALIGN.CENTER
    op0.space_after = Pt(8)

    op1 = otf.add_paragraph()
    op1.text = "DevAssist AI – Smart, Safe, and Automated Code Review"
    op1.font.name = FONT_BODY
    op1.font.size = Pt(15)
    op1.font.bold = True
    op1.font.color.rgb = ACCENT_PRIMARY
    op1.alignment = PP_ALIGN.CENTER
    op1.space_after = Pt(4)

    op2 = otf.add_paragraph()
    op2.text = "We welcome your questions, feedback, and live demonstration!"
    op2.font.name = FONT_BODY
    op2.font.size = Pt(12)
    op2.font.color.rgb = ACCENT_CYAN
    op2.alignment = PP_ALIGN.CENTER

    # 3 Credit Cards
    credits = [
        ("PROJECT CREATORS", "Atharv Jain\nYashvardhan", "AI DevOps Engineering Team", ACCENT_CYAN),
        ("INSTITUTION", "BML Munjal University", "Department of Computer Science", ACCENT_PRIMARY),
        ("FACULTY MENTOR", "Mr. Naveen Singh", "Project Advisor & Guide", ACCENT_EMERALD),
    ]

    cw = Inches(2.8)
    cy = Inches(4.3)
    cgap = Inches(0.4)
    cx_start = Inches(2.06)

    for i, (cr_role, cr_name, cr_sub, cr_acc) in enumerate(credits):
        x = cx_start + i * (cw + cgap)
        cctf = create_card_with_text(s12, x, cy, cw, Inches(1.6), bg_color=CARD_BG_HOVER, accent_bar=cr_acc, pad_x=0.15, pad_y=0.16)

        cp0 = cctf.paragraphs[0]
        cp0.text = cr_role
        cp0.font.name = FONT_BODY
        cp0.font.size = Pt(9.5)
        cp0.font.bold = True
        cp0.font.color.rgb = cr_acc
        cp0.alignment = PP_ALIGN.CENTER
        cp0.space_after = Pt(4)

        cp1 = cctf.add_paragraph()
        cp1.text = cr_name
        cp1.font.name = FONT_TITLE
        cp1.font.size = Pt(12)
        cp1.font.bold = True
        cp1.font.color.rgb = TEXT_WHITE
        cp1.alignment = PP_ALIGN.CENTER
        cp1.space_after = Pt(2)

        cp2 = cctf.add_paragraph()
        cp2.text = cr_sub
        cp2.font.name = FONT_BODY
        cp2.font.size = Pt(8.5)
        cp2.font.color.rgb = TEXT_MUTED
        cp2.alignment = PP_ALIGN.CENTER

    prs.save(out_path)
    print(f"Clean presentation successfully built and saved to: {out_path}")

if __name__ == "__main__":
    out_file = "E:/Devassist-AI-main/devassist-ai/DevAssist_AI_Clean_Presentation.pptx"
    build_presentation(out_file)
