import sys
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def set_shape_text(shape, new_text):
    """Update shape text while preserving font properties of the original text."""
    if not shape.has_text_frame:
        return
    
    tf = shape.text_frame
    
    # Capture original styling from first paragraph/run
    font_name = "Segoe UI"
    font_size = Pt(12)
    font_bold = False
    font_color = None
    
    if tf.paragraphs and tf.paragraphs[0].runs:
        run = tf.paragraphs[0].runs[0]
        if run.font.name:
            font_name = run.font.name
        if run.font.size:
            font_size = run.font.size
        font_bold = run.font.bold
        try:
            if run.font.color and run.font.color.rgb:
                font_color = run.font.color.rgb
        except Exception:
            pass

    # Clear and set new text
    tf.clear()
    lines = new_text.split("\n")
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        
        r = p.add_run()
        r.text = line
        r.font.name = font_name
        r.font.size = font_size
        r.font.bold = font_bold
        if font_color:
            r.font.color.rgb = font_color

def main():
    src_path = "E:/Devassist-AI-main/devassist-ai/AcquiSense_Presentation_Updated.pptx"
    out_path = "E:/Devassist-AI-main/devassist-ai/DevAssist_AI_Presentation.pptx"

    prs = Presentation(src_path)

    # Dictionary of replacements per slide: {slide_index (0-indexed): {shape_index: new_text}}
    replacements = {
        # ── SLIDE 1: Title ────────────────────────────────────────────────────
        0: {
            0: "DevAssist AI – Autonomous Code Review & Architectural Documentation Platform",
            2: "Team Members: Atharv Jain, Yashvardhan",
            4: "University: BML Munjal University",
            6: "Course / Subject: AI DevOps",
            8: "Mentor: Naveen Singh",
        },

        # ── SLIDE 2: Problem → Aim ───────────────────────────────────────────
        1: {
            0: "Problem → Aim: Why DevAssist AI?",
            2: "The problem",
            4: "PR Review Bottlenecks",
            7: "Manual Diff Reading",
            10: "Multi-Language Monorepos",
            13: "Hidden Blast Radius",
            16: "Slow Sprint Turnaround",
            17: "Dozens of PR diffs, cross-file imports, linters, security CVEs, breaking public API changes – reviewed manually under sprint deadlines.",
            18: "The aim",
            20: "DevAssist AI aims to accelerate, automate, and elevate code reviews by using an ensemble multi-LLM review pipeline, AST-aware CodeGraph blast radius calculation, and Hybrid RAG context to identify evidence-backed code defects, generate inline patch fixes, synthesize architectural prologues, and provide real-time dashboard analytics.",
            22: "FAST",
            24: "MULTI-MODEL",
            26: "EVIDENCE-BACKED",
            28: "HUMAN-IN-THE-LOOP",
            30: "Instant first-pass reviews",
            32: "Multi-stage ensemble reasoning",
            34: "Transitive blast-radius detection",
            36: "Confidence scoring & verification",
            38: "Inline Git patch suggestions",
            40: "Final review decision stays with developers",
        },

        # ── SLIDE 3: Multi-Agent / Multi-Stage Review ─────────────────────────
        2: {
            0: "Our Idea: Multi-Agent & Multi-LLM Review System",
            2: "Specialized Review Pipeline & Router",
            7: "Distiller",
            8: "Context & risk digest",
            13: "Reasoner",
            14: "Deep bug & security analysis",
            19: "Validator",
            20: "Confidence scoring & filtering",
            25: "CodeGraph",
            26: "Symbol & caller impact map",
            31: "Hybrid RAG",
            32: "FAISS + BM25 + RRF retrieval",
            37: "Linters",
            38: "Pylint, ESLint, Checkstyle",
            43: "Sandbox",
            44: "Docker agentic test scripts",
            49: "Conversation",
            50: "Contextual PR reply agent",
            55: "Doc Agent",
            56: "Post-merge automated docstrings",
            57: "Instead of a single brittle prompt, DevAssist coordinates specialized stages for distillation, deep reasoning, validation, and impact analysis.",
            58: "How an agent works",
            60: "Diff Patch + CodeGraph\n+ AST Context & Linter Output\n+ Claude 3.5 Sonnet / Gemini",
            63: "Specialized Agent (e.g. Deep Reasoner)",
            66: "Structured Finding:\nLine + Severity + Category + Message + Code Fix",
            67: "LLM = underlying language intelligence.\nAgent = role + instructions + context + LLM + validated structured output.",
        },

        # ── SLIDE 4: End-to-End Workflow ─────────────────────────────────────
        3: {
            0: "How DevAssist AI Works – End-to-End",
            3: "Developer Opens PR",
            6: "Webhook Ingress (HMAC)",
            9: "Redis Queue & Celery",
            12: "CodeGraph & Hybrid RAG",
            15: "Docker Sandbox & Linters",
            18: "3-Stage Ensemble LLM",
            21: "Finding Validator (Gate)",
            24: "Prologue & Chapter Synthesis",
            27: "Inline Comments on GitHub",
            30: "Redis Pub/Sub SSE Stream",
            33: "Next.js 16 Dashboard",
            36: "AST Chunking & CodeGraph",
            37: "Source files parsed by AST & Tree-sitter preserving functions and classes. CodeGraph computes caller dependencies and blast radius up to 3 hops deep.",
            40: "Async Celery & Concurrency",
            41: "Celery background task orchestrates the review asynchronously. Changed files are evaluated concurrently with thread-pool safety to minimize latency.",
            44: "Real-Time SSE Telemetry",
            45: "Celery workers publish live progress events over Redis Pub/Sub channels to FastAPI, streaming instant updates to the Next.js 16 frontend.",
        },

        # ── SLIDE 5: Findings to Decisions ───────────────────────────────────
        4: {
            0: "From Findings to Decisions",
            2: "1  CodeGraph Blast Radius",
            3: "2  Quality Gate (Validator)",
            4: "3  PR Quality Scoring",
            6: "Modified File\n→ utils/auth.py",
            8: "Downstream Caller\n→ api/routes/login.py",
            12: "Impact Analyzer (BFS)",
            15: "Blast Radius: 14 affected callers (HIGH Severity)",
            16: "Transitive risks detected before merging code.",
            21: "Target line in diff hunk?",
            26: "Syntax-valid code fix snippet?",
            31: "Confidence score >= 0.6?",
            36: "Verified against AST context?",
            41: "Duplicate / false-positive check?",
            42: "Raw Findings → Validator → Validated Findings → PR Comment",
            44: "🔴",
            46: "Error (Security / Bug)",
            48: "🟡",
            50: "Warning (Perf / Reliability)",
            52: "🔵",
            54: "Note (Style / Maintainability)",
            56: "Score",
            58: "PR Quality: 8.5 / 10",
            60: "Approve / Comment / Request Changes",
            63: "Human Developer in Control",
            65: "Finding Validator removes hallucinations and filters out noisy, low-confidence suggestions.",
            67: "DevAssist AI provides automated decision support; the final merge decision remains with the engineer.",
        },

        # ── SLIDE 6: Architecture & Tech Stack ────────────────────────────────
        5: {
            0: "System Architecture & Technology Stack",
            3: "Next.js 16 Frontend",
            6: "FastAPI v3.0 REST & Webhook",
            9: "PostgreSQL 15",
            12: "Redis Pub/Sub & Broker",
            16: "Ensemble Pipeline",
            19: "AST Chunker & CodeGraph",
            22: "Celery Worker Queue",
            25: "Distill, Reason & Validate",
            28: "Docker Sandboxed Execution",
            31: "GitHub / GitLab PR Comments",
            36: "Frontend",
            37: "Next.js 16 (Turbopack)\nReact 19\nTailwind CSS v4\n@pierre/diffs\nReactFlow & Mermaid",
            42: "Backend & Queue",
            43: "FastAPI v3.0 (async)\nCelery 5.3+\nRedis 7 (Broker & Pub/Sub)\nPostgreSQL 15 (asyncpg)\nSQLAlchemy 2.0",
            48: "AI & Context Engine",
            49: "Claude 3.5 Sonnet\nGoogle Gemini 2.5 Flash\nOpenAI GPT-4o\nFAISS & rank-bm25 (RRF)\nTree-sitter & AST",
            54: "Infrastructure & Security",
            55: "Docker & Docker Compose\nDocker Sandbox (Network-none)\nPyngrok Tunneling\nGitHub App / Webhooks\nClerk Authentication",
        },

        # ── SLIDE 7: What User Sees ──────────────────────────────────────────
        6: {
            0: "What the User Sees: Developer Dashboard & PR Comments",
            5: "Engineering Overview & PR Hub",
            8: "Review Velocity",
            20: "Severity Breakdown",
            39: "PR Health Score",
            41: "8.5 / 10",
            44: "Pierre Diff Viewer",
            50: "ReactFlow Blast Radius Graph",
            58: "Expected impact (project goals)",
            61: "70%+ faster review turnaround times",
            64: "Interactive downstream caller blast-radius maps",
            67: "Evidence-backed findings with suggested Git fixes",
            70: "Chapter-clustered diff walkthroughs for big PRs",
            73: "Real-time SSE progress streaming during reviews",
            76: "Automatic architectural Mermaid diagram synthesis",
            79: "Developer retains full control of code approvals",
        },

        # ── SLIDE 8: Where AI Reduces Effort ─────────────────────────────────
        7: {
            0: "Where AI Reduces Effort and Improves Code Review",
            3: "Traditional manual workflow",
            5: "DevAssist AI-assisted workflow",
            7: "Manual line-by-line diff inspection",
            10: "AST-aware hybrid context & chunking",
            12: "Overlooked downstream caller impacts",
            15: "Automated CodeGraph BFS blast-radius map",
            17: "Manual linter execution & syntax debates",
            20: "Parallel static linters & Docker sandbox",
            22: "Unfiltered single-LLM hallucinations",
            25: "3-Stage Distill-Reason-Validate pipeline",
            27: "Forgotten docstrings and release notes",
            30: "Autonomous post-merge docstring generation",
            32: "Manual, time-consuming, prone to human fatigue and missed regressions.",
            34: "Automated diff enrichment, multi-model reasoning, verified code fixes, real-time telemetry, and PR walkthroughs.",
        },

        # ── SLIDE 9: What Improves ───────────────────────────────────────────
        8: {
            0: "What Improves with DevAssist AI?",
            6: "Review Speed",
            7: "Async Celery execution and parallel file evaluation cut PR turnaround time by over 70%.",
            12: "Blast Radius",
            13: "CodeGraph maps transitive callers up to 3 hops deep to prevent breaking public API changes.",
            18: "Low Hallucination",
            19: "Stage 3 Finding Validator prunes ungrounded findings below 0.6 confidence threshold.",
            24: "Multi-LLM Uptime",
            25: "Fallback router seamlessly transitions across OpenAI, Claude, Gemini, NVIDIA, and Ollama.",
            30: "Developer UX",
            31: "One-click suggested Git patches, Mermaid architecture diagrams, and conversational PR replies.",
        },

        # ── SLIDE 10: DevOps & Deployment ────────────────────────────────────
        9: {
            0: "DevOps & Deployment: From Development to Monitoring",
            3: "Development / Code",
            6: "Containerization (Docker)",
            9: "Service Config (Compose)",
            12: "Tunneling (Pyngrok)",
            15: "Docker Security Sandbox",
            18: "SSE Telemetry (Redis)",
            21: "Health & Observability",
            22: "Supports the full containerized stack:",
            24: "Next.js 16 Frontend",
            26: "FastAPI v3.0 API Server",
            28: "Celery Async Workers",
            30: "PostgreSQL 15 (asyncpg)",
            32: "Redis 7 (Broker & Pub/Sub)",
            35: "Docker & Compose",
            36: "Docker containerizes Redis, PostgreSQL, and Celery workers for 100% reproducible environments. Compose orchestrates dependencies with active healthchecks.",
            39: "Isolated Dynamic Sandbox",
            40: "AgenticExecutor runs AI-generated validation scripts in network-isolated, read-only Docker containers as an unprivileged user with dropped capabilities.",
            43: "Operational Telemetry",
            44: "Real-time Server-Sent Events over Redis Pub/Sub, live /health endpoint monitoring, and automated SQLAlchemy migrations via Alembic.",
        },

        # ── SLIDE 11: Future Scope ───────────────────────────────────────────
        10: {
            0: "Future Scope",
            6: "IDE Extensions",
            7: "Live pre-commit code review inside VS Code and JetBrains IDEs.",
            12: "Self-Healing PRs",
            13: "Automated branch creation with verified fix commits and passing test runs.",
            18: "Cross-Repo Graphs",
            19: "Multi-repository and microservice contract dependency blast-radius mapping.",
            24: "Fine-Tuned Models",
            25: "Custom fine-tuned on-premise models trained on corporate coding standards.",
            30: "Extended Linters",
            31: "Semgrep, SonarQube, and OWASP Dependency-Check deep integrations.",
            36: "Enterprise Governance",
            37: "SAML SSO, custom review policies, and immutable SOC2-compliant audit trails.",
            42: "Automated Test Gen",
            43: "Autonomous generation of regression unit tests for new pull request diffs.",
            45: "Goal: make DevAssist AI the most reliable, context-aware, and secure autonomous engineering assistant for modern development teams.",
        },

        # ── SLIDE 12: Thank You ──────────────────────────────────────────────
        11: {
            0: "Thank You",
            2: "DevAssist AI – Autonomous Code Review & Architectural Documentation Platform\nTeam Members: Atharv Jain, Yashvardhan | BML Munjal University | Mentor: Naveen Singh",
        }
    }

    # Execute modifications
    for slide_idx, slide_replacements in replacements.items():
        if slide_idx >= len(prs.slides):
            continue
        slide = prs.slides[slide_idx]
        for shape_idx, new_text in slide_replacements.items():
            if shape_idx < len(slide.shapes):
                shape = slide.shapes[shape_idx]
                if shape.has_text_frame:
                    set_shape_text(shape, new_text)

    prs.save(out_path)
    print(f"Presentation successfully generated and saved to: {out_path}")

if __name__ == "__main__":
    main()
